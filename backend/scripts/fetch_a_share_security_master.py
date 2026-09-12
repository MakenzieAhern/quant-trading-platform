"""Fetch a pilot-complete A-share security master from official exchange surfaces."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import time
import urllib.parse
import urllib.request
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

PILOT_START = date(2023, 1, 1)
PILOT_END = date(2024, 7, 31)
RETRIEVAL_END = date.today()
PARSER_VERSION = "official-a-share-pilot-v2-a-only-exits"
SSE_PAGE = "https://www.sse.com.cn/assortment/stock/list/share/"
SSE_DELIST_PAGE = "https://www.sse.com.cn/assortment/stock/list/delisting/"
SZSE_LIST_PAGE = "https://www.szse.cn/market/product/stock/list/index.html"
SZSE_NOTICE_PAGE = "https://www.szse.cn/disclosure/notice/company/index.html"
SZSE_DISCLOSURE_PAGE = "https://www.szse.cn/disclosure/listed/notice/index.html"
SZSE_YEARBOOK_PAGE = "https://www.szse.cn/market/periodical/year/index.html"
SZSE_YEARBOOK_URL = (
    "https://docs.static.szse.cn/www/market/periodical/year/W020230915392900565815.pdf"
)
FIELDS = (
    "security_id",
    "exchange",
    "symbol",
    "name",
    "board",
    "listing_date",
    "delisting_date",
    "effective_from",
    "effective_to",
    "status",
    "source_url",
)


class FetchError(RuntimeError):
    pass


def _json_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def _sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class OfficialFetcher:
    def __init__(
        self,
        root: Path,
        *,
        retrieved_at: str | None = None,
        cached: dict[str, dict[str, object]] | None = None,
    ) -> None:
        self.root = root
        self.retrieved_at = retrieved_at or datetime.now(UTC).isoformat().replace("+00:00", "Z")
        self.cached = cached or {}
        self.snapshots: list[dict[str, object]] = []

    def get(self, name: str, url: str, *, exchange: str, referer: str) -> bytes:
        return self._request(name, url, exchange=exchange, referer=referer)

    def post(
        self,
        name: str,
        url: str,
        data: dict[str, object],
        *,
        exchange: str,
        referer: str,
        json_body: bool = False,
    ) -> bytes:
        if json_body:
            body = json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode()
            content_type = "application/json"
        else:
            body = urllib.parse.urlencode(data, doseq=True).encode()
            content_type = "application/x-www-form-urlencoded"
        return self._request(
            name,
            url,
            exchange=exchange,
            referer=referer,
            body=body,
            content_type=content_type,
        )

    def _request(
        self,
        name: str,
        url: str,
        *,
        exchange: str,
        referer: str,
        body: bytes | None = None,
        content_type: str | None = None,
    ) -> bytes:
        cached = self.cached.get(name)
        cached_path = self.root / name
        if (
            cached
            and cached.get("source_url") == url
            and cached_path.is_file()
            and _sha256_file(cached_path) == cached.get("sha256")
        ):
            self.snapshots.append(dict(cached))
            return cached_path.read_bytes()
        headers = {"User-Agent": "Mozilla/5.0", "Referer": referer}
        if content_type:
            headers["Content-Type"] = content_type
        request = urllib.request.Request(url, data=body, headers=headers)
        last_error: Exception | None = None
        for attempt in range(3):
            try:
                with urllib.request.urlopen(request, timeout=60) as response:
                    payload = response.read()
                    response_headers = dict(response.headers.items())
                    status = response.status
                    final_url = response.url
                break
            except (OSError, TimeoutError) as exc:
                last_error = exc
                if attempt == 2:
                    raise FetchError(f"OFFICIAL_SOURCE_FETCH_FAILED:{name}") from exc
                time.sleep(1 + attempt)
        else:  # pragma: no cover
            raise FetchError(f"OFFICIAL_SOURCE_FETCH_FAILED:{name}") from last_error
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        self.snapshots.append(
            {
                "exchange": exchange,
                "source_url": url,
                "final_url": final_url,
                "retrieved_at": self.retrieved_at,
                "http_status": status,
                "http_headers": response_headers,
                "sha256": hashlib.sha256(payload).hexdigest(),
                "parser_version": PARSER_VERSION,
                "raw_snapshot_filename": name,
            }
        )
        time.sleep(0.35)
        return payload


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-root",
        type=Path,
        default=Path("data/research/staging/security-master-pilot"),
    )
    parser.add_argument("--reuse-snapshots", action="store_true")
    args = parser.parse_args()
    _build(args.output_root, reuse_snapshots=args.reuse_snapshots)


def _build(root: Path, *, reuse_snapshots: bool = False) -> None:
    root.mkdir(parents=True, exist_ok=True)
    prior = None
    snapshot_path = root / "source-snapshot-manifest.json"
    if reuse_snapshots and snapshot_path.is_file():
        prior = json.loads(snapshot_path.read_text(encoding="utf-8"))
    fetcher = OfficialFetcher(
        root / "raw",
        retrieved_at=prior["retrieved_at"] if prior else None,
        cached={item["raw_snapshot_filename"]: item for item in prior["snapshots"]}
        if prior
        else None,
    )
    sse_rows = _fetch_sse(fetcher)
    szse_rows, reconciliation = _fetch_szse(fetcher)
    conflicts = _conflicts((*sse_rows, *szse_rows))
    complete = not conflicts and not reconciliation["unresolved_decisions"]
    normalized = {
        "SSE_OFFICIAL": root / "sse-security-master.csv",
        "SZSE_OFFICIAL": root / "szse-security-master.csv",
    }
    for source, rows in (("SSE_OFFICIAL", sse_rows), ("SZSE_OFFICIAL", szse_rows)):
        _write_csv(normalized[source], rows)
        checksum = _sha256_file(normalized[source])
        for snapshot in fetcher.snapshots:
            if snapshot["exchange"] == source.split("_")[0]:
                snapshot["normalized_output_checksum"] = checksum
    manifest = {
        "schema_version": "1.0",
        "coverage_start": PILOT_START.isoformat(),
        "coverage_end": PILOT_END.isoformat(),
        "complete": complete,
        "sources": [
            {
                "source": source,
                "source_version": fetcher.retrieved_at[:10],
                "filename": path.name,
                "sha256": _sha256_file(path),
                "source_url": SSE_PAGE if source == "SSE_OFFICIAL" else SZSE_LIST_PAGE,
                "retrieval_timestamp": fetcher.retrieved_at,
                "complete": complete,
            }
            for source, path in normalized.items()
        ],
    }
    (root / "security-master-input.json").write_bytes(_json_bytes(manifest) + b"\n")
    snapshot_manifest = {
        "schema_version": "1.0",
        "retrieved_at": fetcher.retrieved_at,
        "parser_version": PARSER_VERSION,
        "pilot_start": PILOT_START.isoformat(),
        "pilot_end": PILOT_END.isoformat(),
        "complete": complete,
        "conflicts": conflicts,
        "reconciliation": reconciliation,
        "snapshots": sorted(fetcher.snapshots, key=lambda item: str(item["raw_snapshot_filename"])),
        "checksum": hashlib.sha256(
            _json_bytes(
                {
                    "normalized": {key: _sha256_file(path) for key, path in normalized.items()},
                    "snapshots": [item["sha256"] for item in fetcher.snapshots],
                }
            )
        ).hexdigest(),
    }
    (root / "source-snapshot-manifest.json").write_bytes(_json_bytes(snapshot_manifest) + b"\n")
    print(
        json.dumps(
            {key: value for key, value in snapshot_manifest.items() if key != "snapshots"},
            ensure_ascii=True,
            sort_keys=True,
        )
    )
    if not complete:
        raise SystemExit(2)


def _fetch_sse(fetcher: OfficialFetcher) -> list[dict[str, str]]:
    active: dict[str, dict[str, str]] = {}
    for stock_type, board in (("1", "SSE_MAIN"), ("8", "SSE_STAR")):
        params = {
            "sqlId": "COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L",
            "isPagination": "true",
            "STOCK_CODE": "",
            "CSRC_CODE": "",
            "REG_PROVINCE": "",
            "STOCK_TYPE": stock_type,
            "COMPANY_STATUS": "2,4,5,7,8",
            "type": "inParams",
            "pageHelp.cacheSize": 1,
            "pageHelp.beginPage": 1,
            "pageHelp.pageSize": 2000,
            "pageHelp.pageNo": 1,
        }
        url = "https://query.sse.com.cn/sseQuery/commonQuery.do?" + urllib.parse.urlencode(params)
        payload = fetcher.get(
            f"sse-current-{stock_type}.json", url, exchange="SSE", referer=SSE_PAGE
        )
        for item in _sse_data(payload):
            symbol = str(item["A_STOCK_CODE"])
            listing = _compact_date(str(item["LIST_DATE"]))
            if listing <= PILOT_END:
                active[symbol] = _row(
                    "SSE", symbol, str(item["COMPANY_ABBR"]), board, listing, None, SSE_PAGE
                )
    params["sqlId"] = "COMMON_SSE_CP_GPJCTPZ_GPLB_ZZGP_L"
    # Type 2 is B shares: its LIST_DATE can differ for the same A_STOCK_CODE.
    params["STOCK_TYPE"] = "1,8"
    params["COMPANY_STATUS"] = "3"
    url = "https://query.sse.com.cn/sseQuery/commonQuery.do?" + urllib.parse.urlencode(params)
    payload = fetcher.get(
        "sse-terminated-a-boards.json", url, exchange="SSE", referer=SSE_DELIST_PAGE
    )
    for item in _sse_data(payload):
        symbol = str(item["A_STOCK_CODE"])
        if not re.fullmatch(r"6\d{5}", symbol):
            continue
        listing = _compact_date(str(item["LIST_DATE"]))
        delisting = _compact_date(str(item["DELIST_DATE"]))
        if listing <= PILOT_END and delisting >= PILOT_START:
            board = "SSE_STAR" if symbol.startswith(("688", "689")) else "SSE_MAIN"
            active[symbol] = _row(
                "SSE",
                symbol,
                str(item["COMPANY_ABBR"]),
                board,
                listing,
                delisting,
                SSE_DELIST_PAGE,
            )
    return sorted(active.values(), key=lambda item: item["security_id"])


def _fetch_szse(fetcher: OfficialFetcher) -> tuple[list[dict[str, str]], dict[str, object]]:
    list_url = "https://www.szse.cn/api/report/ShowReport?SHOWTYPE=xlsx&CATALOGID=1110&TABKEY=tab1"
    payload = fetcher.get("szse-current.xlsx", list_url, exchange="SZSE", referer=SZSE_LIST_PAGE)
    active: dict[str, dict[str, str]] = {}
    sheet = load_workbook(io.BytesIO(payload), read_only=False, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = tuple(str(value) for value in next(rows))
    for values in rows:
        item = dict(zip(header, values, strict=True))
        symbol = str(item.get("A股代码") or "")
        if not _szse_a_symbol(symbol):
            continue
        listing = date.fromisoformat(str(item["A股上市日期"]))
        if listing <= PILOT_END:
            board = "SZSE_CHINEXT" if symbol.startswith(("300", "301")) else "SZSE_MAIN"
            active[symbol] = _row(
                "SZSE", symbol, str(item["A股简称"]), board, listing, None, SZSE_LIST_PAGE
            )
    decisions = _exchange_notices(fetcher)
    final_notices = _final_delisting_notices(fetcher)
    profiles: dict[str, dict[str, Any]] = {}
    delisted: dict[str, dict[str, str]] = {}
    for item in final_notices:
        symbol = next((str(code) for code in item["secCode"] if _szse_a_symbol(str(code))), "")
        if not symbol:
            continue
        profile_url = (
            "https://www.szse.cn/api/report/index/companyGeneralization?"
            + urllib.parse.urlencode({"secCode": symbol})
        )
        profile_payload = fetcher.get(
            f"szse-profile-{symbol}.json",
            profile_url,
            exchange="SZSE",
            referer=f"https://www.szse.cn/certificate/individual/index.html?code={symbol}",
        )
        profile = json.loads(profile_payload)
        if str(profile.get("code")) != "0" or not profile.get("data"):
            continue
        profiles[symbol] = profile
        attachment = str(item["attachPath"])
        pdf_url = "https://disc.static.szse.cn/download" + attachment
        pdf = fetcher.get(
            f"szse-final-delist-{symbol}.pdf",
            pdf_url,
            exchange="SZSE",
            referer=SZSE_DISCLOSURE_PAGE,
        )
        delisting = extract_delisting_date(pdf)
        listing = date.fromisoformat(str(profile["data"]["agssrq"]))
        if listing <= PILOT_END and delisting >= PILOT_START:
            board = "SZSE_CHINEXT" if symbol.startswith(("300", "301")) else "SZSE_MAIN"
            delisted[symbol] = _row(
                "SZSE",
                symbol,
                str(profile["data"]["agjc"]),
                board,
                listing,
                delisting,
                pdf_url,
            )
    # A rare missing issuer notice is published directly by SZSE as an exchange notice.
    for decision in decisions:
        if "暨摘牌" not in decision["title"]:
            continue
        html = fetcher.get(
            f"szse-exchange-final-{decision['slug']}.html",
            decision["url"],
            exchange="SZSE",
            referer=SZSE_NOTICE_PAGE,
        ).decode("utf-8", errors="replace")
        symbol_match = re.search(r"股票代码\s*([0-9]{6})", _html_text(html))
        date_match = re.search(
            r"摘牌日期\s*[\uff1a:]\s*([0-9]{4}年[0-9]+月[0-9]+日)",
            _html_text(html),
        )
        if not symbol_match or not date_match:
            continue
        symbol = symbol_match.group(1)
        profile_url = (
            "https://www.szse.cn/api/report/index/companyGeneralization?"
            + urllib.parse.urlencode({"secCode": symbol})
        )
        profile = json.loads(
            fetcher.get(
                f"szse-profile-{symbol}.json",
                profile_url,
                exchange="SZSE",
                referer=f"https://www.szse.cn/certificate/individual/index.html?code={symbol}",
            )
        )
        profiles[symbol] = profile
        listing = date.fromisoformat(str(profile["data"]["agssrq"]))
        delisting = _chinese_date(date_match.group(1))
        board = "SZSE_CHINEXT" if symbol.startswith(("300", "301")) else "SZSE_MAIN"
        delisted[symbol] = _row(
            "SZSE",
            symbol,
            str(profile["data"]["agjc"]),
            board,
            listing,
            delisting,
            decision["url"],
        )
    active.update(delisted)
    baseline, b_share_lines = _yearbook_baseline(fetcher)
    reconstructed = {
        symbol
        for symbol, item in active.items()
        if date.fromisoformat(item["listing_date"]) <= date(2022, 12, 31)
        and (
            not item["delisting_date"]
            or date.fromisoformat(item["delisting_date"]) > date(2022, 12, 31)
        )
    }
    for symbol in sorted(baseline - reconstructed):
        profile_url = (
            "https://www.szse.cn/api/report/index/companyGeneralization?"
            + urllib.parse.urlencode({"secCode": symbol})
        )
        profile = json.loads(
            fetcher.get(
                f"szse-profile-{symbol}.json",
                profile_url,
                exchange="SZSE",
                referer=f"https://www.szse.cn/certificate/individual/index.html?code={symbol}",
            )
        )
        if not (profile.get("data") or {}).get("agdm"):
            continue
        profiles[symbol] = profile
        company = _name_key(str(profile["data"]["gsqc"]))
        decision = next((item for item in decisions if _name_key(item["company"]) == company), None)
        delisting = None
        source_url = profile_url
        if decision:
            html = fetcher.get(
                f"szse-exchange-termination-{decision['slug']}.html",
                decision["url"],
                exchange="SZSE",
                referer=SZSE_NOTICE_PAGE,
            ).decode("utf-8", errors="replace")
            match = re.search(
                r"股票自\s*([0-9]{4}\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日)\s*起终止上市",
                _html_text(html),
            )
            if not match:
                continue
            delisting = _chinese_date(match.group(1))
            source_url = decision["url"]
        listing = date.fromisoformat(str(profile["data"]["agssrq"]))
        active[symbol] = _row(
            "SZSE",
            symbol,
            str(profile["data"]["agjc"]),
            "SZSE_CHINEXT" if symbol.startswith(("300", "301")) else "SZSE_MAIN",
            listing,
            delisting,
            source_url,
        )
    represented_names = {_name_key(str(profile["data"]["gsqc"])) for profile in profiles.values()}
    unresolved = sorted(
        decision["title"]
        for decision in decisions
        if _name_key(decision["company"]) not in represented_names
        and not any(_name_key(decision["company"]) in line for line in b_share_lines)
    )
    reconstructed = {
        symbol
        for symbol, item in active.items()
        if date.fromisoformat(item["listing_date"]) <= date(2022, 12, 31)
        and (
            not item["delisting_date"]
            or date.fromisoformat(item["delisting_date"]) > date(2022, 12, 31)
        )
    }
    baseline_missing = sorted(baseline - reconstructed)
    baseline_extra = sorted(reconstructed - baseline)
    if baseline_missing or baseline_extra:
        unresolved.append(
            f"SZSE_2022_BASELINE_MISMATCH:missing={len(baseline_missing)}:extra={len(baseline_extra)}"
        )
    return sorted(active.values(), key=lambda item: item["security_id"]), {
        "exchange_termination_decisions": len(decisions),
        "final_delisting_notices": len(final_notices),
        "yearbook_2022_a_share_codes": len(baseline),
        "reconstructed_2022_a_share_codes": len(reconstructed),
        "baseline_missing": baseline_missing,
        "baseline_extra": baseline_extra,
        "unresolved_decisions": unresolved,
    }


def _exchange_notices(fetcher: OfficialFetcher) -> list[dict[str, str]]:
    start_ms = int(datetime(2023, 1, 1, tzinfo=UTC).timestamp() * 1000)
    end_ms = int(
        datetime.combine(RETRIEVAL_END, datetime.max.time(), tzinfo=UTC).timestamp() * 1000
    )
    found: list[dict[str, str]] = []
    page = 1
    while True:
        payload = fetcher.post(
            f"szse-exchange-notices-{page}.json",
            "https://www.szse.cn/api/search/content",
            {
                "keyword": "",
                "time": start_ms,
                "endTime": end_ms,
                "range": "title",
                "channelCode[]": "ipocompany_news",
                "currentPage": page,
                "pageSize": 50,
                "scope": 1,
            },
            exchange="SZSE",
            referer=SZSE_NOTICE_PAGE,
        )
        result = json.loads(payload)
        for item in result["data"]:
            title = re.sub(r"<[^>]+>", "", str(item["doctitle"]))
            match = re.fullmatch(r"关于(.+)股票终止上市(暨摘牌)?的公告", title)
            if match:
                url = str(item["docpuburl"]).replace("http://", "https://")
                found.append(
                    {
                        "title": title,
                        "company": match.group(1),
                        "url": url,
                        "slug": Path(urllib.parse.urlparse(url).path).stem,
                    }
                )
        if page * int(result["pageSize"]) >= int(result["totalSize"]):
            break
        page += 1
    unique = {item["title"]: item for item in found}
    return sorted(unique.values(), key=lambda item: item["title"])


def _final_delisting_notices(fetcher: OfficialFetcher) -> list[dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    page = 1
    while True:
        payload = fetcher.post(
            f"szse-final-delist-notices-{page}.json",
            "https://www.szse.cn/api/disc/announcement/annList",
            {
                "pageSize": 50,
                "pageNum": page,
                "searchKey": ["终止上市暨摘牌"],
                "channelCode": ["listedNotice_disc"],
                "seDate": [PILOT_START.isoformat(), RETRIEVAL_END.isoformat()],
            },
            exchange="SZSE",
            referer=SZSE_DISCLOSURE_PAGE,
            json_body=True,
        )
        result = json.loads(payload)
        for item in result["data"]:
            title = str(item["title"])
            if re.search(r"股票.*终止上市.*摘牌|股票终止上市暨摘牌", title):
                for code in item.get("secCode", []):
                    if _szse_a_symbol(str(code)):
                        found[str(code)] = item
        if page * 50 >= int(result["announceCount"]):
            break
        page += 1
    return [found[key] for key in sorted(found)]


def _yearbook_baseline(fetcher: OfficialFetcher) -> tuple[set[str], set[str]]:
    pdf = fetcher.get(
        "szse-yearbook-2022.pdf",
        SZSE_YEARBOOK_URL,
        exchange="SZSE",
        referer=SZSE_YEARBOOK_PAGE,
    )
    reader = PdfReader(io.BytesIO(pdf))
    collecting = False
    codes: set[str] = set()
    b_share_lines: set[str] = set()
    for page in reader.pages[-160:]:
        text = page.extract_text() or ""
        if "NAMELIST OF LISTED COMPANIES" in text:
            collecting = True
        if collecting and (
            "ADDRESS LIST OF LISTED COMPANIES" in text or "公司全称\nCOMPANY\n地址\nADDRESS" in text
        ):
            break
        if collecting:
            for code, short_name, line in re.findall(r"(?m)^([0-9]{6})\s+(\S+)\s+(.+)$", text):
                if _factbook_b_share(code, short_name, line):
                    b_share_lines.add(_name_key(line))
                elif _szse_a_symbol(code):
                    codes.add(code)
    if len(codes) < 2000:
        raise FetchError("SZSE_YEARBOOK_BASELINE_PARSE_FAILED")
    return codes, b_share_lines


def _factbook_b_share(code: str, short_name: str, line: str) -> bool:
    return bool(
        re.match(
            r"^[0-9]{6}\s+.+[B\uff22]\s+[\u4e00-\u9fff]",
            f"{code} {short_name} {line}",
            re.IGNORECASE,
        )
    )


def extract_delisting_date(pdf: bytes) -> date:
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
    return _delisting_date_from_text(text)


def _delisting_date_from_text(text: str) -> date:
    patterns = (
        r"摘牌日期\s*[\uff1a:]\s*([0-9]{4}\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日)",
        r"(?:将于|将在|于)\s*([0-9]{4}\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日)\s*(?:被)?摘牌",
        r"终止上市(?:及摘牌)?日期\s*[\uff1a:]\s*([0-9]{4}\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日)",
    )
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return _chinese_date(match.group(1))
    raise FetchError("SZSE_DELISTING_DATE_MISSING")


def _sse_data(payload: bytes) -> list[dict[str, object]]:
    value = json.loads(payload)
    rows = value.get("result") or value.get("pageHelp", {}).get("data")
    if not isinstance(rows, list):
        raise FetchError("SSE_RESPONSE_SCHEMA_INVALID")
    return rows


def _row(
    exchange: str,
    symbol: str,
    name: str,
    board: str,
    listing: date,
    delisting: date | None,
    source_url: str,
) -> dict[str, str]:
    return {
        "security_id": f"CN-A:{exchange}:{symbol}",
        "exchange": exchange,
        "symbol": symbol,
        "name": name.strip(),
        "board": board,
        "listing_date": listing.isoformat(),
        "delisting_date": delisting.isoformat() if delisting else "",
        "effective_from": listing.isoformat(),
        "effective_to": delisting.isoformat() if delisting else "",
        "status": "DELISTED" if delisting else "LISTED",
        "source_url": source_url,
    }


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows({key: row[key] for key in FIELDS} for row in rows)


def _conflicts(rows: tuple[dict[str, str], ...]) -> list[str]:
    grouped: dict[str, set[tuple[str, ...]]] = {}
    for row in rows:
        grouped.setdefault(row["security_id"], set()).add(tuple(row[key] for key in FIELDS))
    return sorted(key for key, values in grouped.items() if len(values) > 1)


def _compact_date(value: str) -> date:
    return datetime.strptime(value, "%Y%m%d").date()


def _chinese_date(value: str) -> date:
    match = re.fullmatch(r"\s*([0-9]{4})\s*年\s*([0-9]{1,2})\s*月\s*([0-9]{1,2})\s*日\s*", value)
    if not match:
        raise FetchError("CHINESE_DATE_INVALID")
    return date(*(int(part) for part in match.groups()))


def _szse_a_symbol(symbol: str) -> bool:
    return bool(re.fullmatch(r"(?:000|001|002|003|300|301)[0-9]{3}", symbol))


def _name_key(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", value).replace("股份有限公司", "")


def _html_text(value: str) -> str:
    return re.sub(r"\s+", "", re.sub(r"<[^>]+>", " ", value)).replace("&nbsp;", "")


if __name__ == "__main__":
    main()
