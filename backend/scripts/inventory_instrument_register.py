"""Verify explicit exchange register types without backdating current snapshots."""

import argparse
import json
import time
from collections import Counter
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

from app.domain.markets.identifiers import InstrumentType
from scripts.reconcile_sse_cninfo import fetch, save, sha
from scripts.stage_akshare_pilot import candidate_members


def sse_type(record: dict) -> InstrumentType:
    return {
        "主板A": InstrumentType.COMMON_A_SHARE,
        "科创A": InstrumentType.COMMON_A_SHARE,
        "科创CDR": InstrumentType.CDR,
        "主板B": InstrumentType.B_SHARE,
    }[record["SEC_TYPE"]]


def build(staging: Path) -> None:
    root = staging / "security-master-pilot"
    output = staging / "instrument-scope-20260912"
    manifest = json.loads((root / "source-snapshot-manifest.json").read_bytes())
    sources = {r["raw_snapshot_filename"]: r for r in manifest["snapshots"]}

    def source(name: str) -> tuple[Path, dict]:
        path, meta = root / "raw" / name, sources[name]
        if sha(path.read_bytes()) != meta["sha256"]:
            raise ValueError("ACCEPTED_REGISTER_CHECKSUM_CHANGED")
        return path, dict(
            path=path.relative_to(staging).as_posix(),
            source_url=meta["source_url"],
            sha256=meta["sha256"],
            retrieval_timestamp_utc=meta["retrieved_at"],
        )

    sse = {}
    for name in ("sse-current-1.json", "sse-current-8.json", "sse-terminated.json"):
        path, meta = source(name)
        for row in json.loads(path.read_bytes())["result"]:
            sse[str(row["A_STOCK_CODE"])] = row, meta
    path, szse_meta = source("szse-current.xlsx")
    # Official workbook dimension metadata says A1 although the sheet has many columns.
    # Match the accepted master parser: normal loading discovers the actual cells.
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    header = next(values)
    szse = {}
    for values_row in values:
        row = dict(zip(header, values_row, strict=True))
        if row.get("A股代码"):
            szse[str(row["A股代码"])] = row
    workbook.close()
    rows = []
    members = candidate_members(staging)
    checkpoint = output / "current-register-inventory.json"
    prior = (
        {r["security_id"]: r for r in json.loads(checkpoint.read_bytes())}
        if checkpoint.exists()
        else {}
    )
    for member in sorted(members, key=lambda r: r["security_id"]):
        symbol, sid = member["symbol"], member["security_id"]
        if sid in prior:
            result = prior[sid]
            if any(
                sha((staging / s["path"]).read_bytes()) != s["sha256"] for s in result["sources"]
            ):
                raise ValueError("INSTRUMENT_REGISTER_CHECKPOINT_CHANGED")
            rows.append(result)
            continue
        try:
            if member["exchange"] == "SSE":
                raw, meta = sse[symbol]
                evidence = [meta]
                if raw["LIST_DATE"] != member["listing_date"].replace("-", ""):
                    raise ValueError("REGISTER_LISTING_IDENTITY_CHANGED")
                if member["board"] == "SSE_STAR":
                    url = "https://query.sse.com.cn/commonQuery.do?" + urlencode(
                        dict(
                            isPagination="false",
                            sqlId="COMMON_SSE_CP_GPJCTPZ_GPLB_GPGK_GSGK_C",
                            COMPANY_CODE=symbol,
                        )
                    )
                    name = f"sse-company-overview-{symbol}.json"
                    existed = (output / name).exists()
                    payload = json.loads(
                        fetch(output, name, url, referer="https://www.sse.com.cn/")
                    )
                    matches = payload["result"]
                    if len(matches) != 1 or matches[0]["A_STOCK_CODE"] != symbol:
                        raise ValueError("REGISTER_SECURITY_IDENTITY_CHANGED")
                    detail = matches[0]
                    if detail["A_LIST_DATE"] != raw["LIST_DATE"]:
                        raise ValueError("OVERVIEW_LISTING_DATE_CHANGED")
                    kind = sse_type(detail)
                    new_meta = json.loads((output / (name + ".meta.json")).read_bytes())
                    evidence.append(
                        dict(
                            path=(output / name).relative_to(staging).as_posix(),
                            source_url=url,
                            sha256=new_meta["sha256"],
                            retrieval_timestamp_utc=new_meta["retrieved_at_utc"],
                        )
                    )
                    fields = {k: detail[k] for k in ("A_STOCK_CODE", "A_LIST_DATE", "SEC_TYPE")}
                    if not existed:
                        time.sleep(0.5)
                else:
                    # Explicit A-security column, not the symbol's numeric prefix.
                    if raw["A_STOCK_CODE"] != symbol or raw["LIST_BOARD"] != "1":
                        raise ValueError("EXPLICIT_MAIN_A_REGISTER_FIELD_MISSING")
                    kind = InstrumentType.COMMON_A_SHARE
                    fields = {
                        k: raw[k]
                        for k in ("A_STOCK_CODE", "B_STOCK_CODE", "LIST_DATE", "LIST_BOARD")
                    }
            else:
                if symbol in szse:
                    raw, evidence = szse[symbol], [szse_meta]
                    fields = {k: raw[k] for k in ("A股代码", "A股上市日期", "A股简称")}
                    listing = str(raw["A股上市日期"])
                else:
                    path, meta = source(f"szse-profile-{symbol}.json")
                    payload = json.loads(path.read_bytes())
                    raw, evidence = payload["data"], [meta]
                    if raw["agdm"] != symbol or payload["cols"]["agdm"] != "A股代码":
                        raise ValueError("EXPLICIT_SZSE_A_REGISTER_FIELD_MISSING")
                    fields = {k: raw[k] for k in ("agdm", "agssrq", "agjc")}
                    listing = raw["agssrq"]
                if listing != member["listing_date"]:
                    raise ValueError("REGISTER_LISTING_DATE_CHANGED")
                kind = InstrumentType.COMMON_A_SHARE
            rows.append(
                dict(
                    **member,
                    instrument_type=kind.value,
                    sources=evidence,
                    source_fields=fields,
                    temporal_status="HISTORICAL_EFFECTIVITY_NOT_YET_CERTIFIED",
                )
            )
            save(checkpoint, rows)
            if len(rows) % 25 == 0 or len(rows) == len(members):
                print(f"classified_current_register={len(rows)}/{len(members)}", flush=True)
        except Exception:
            save(checkpoint, rows)
            raise
    report = dict(
        total_master_securities=len(members),
        classified_current_register=len(rows),
        counts=dict(sorted(Counter(r["instrument_type"] for r in rows).items())),
        unclassified=len(members) - len(rows),
        nonordinary=[r for r in rows if r["instrument_type"] != "COMMON_A_SHARE"],
        historical_effectivity="NOT_YET_CERTIFIED",
        pilot_publication_authorized=False,
        inventory_sha256=sha(checkpoint.read_bytes()),
    )
    save(output / "current-register-audit.json", report)
    print(json.dumps(report, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--staging", required=True, type=Path)
    build(parser.parse_args().staging)
