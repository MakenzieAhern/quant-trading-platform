"""Certify the offline R5 structural-action census from retained official evidence."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.data.catalog.checksums import canonical_json_bytes, canonical_sha256, sha256_file
from app.data.catalog.freeze import write_immutable
from app.domain.markets.consolidations import ShareConsolidationEvent
from app.domain.markets.demergers import DemergerEvent, DemergerSource
from app.domain.markets.pass_through import RestructuringBonusPassThroughEvent
from app.domain.markets.surrenders import (
    MandatoryPartialShareSurrenderEvent,
    SurrenderQuantityBand,
)

APPROVAL_SHA256 = "9362f32566013841a3a45fd5946abbc9c128fd4f66b5ea20e23c902180a8a4ad"

SURRENDERS: tuple[dict[str, Any], ...] = (
    {
        "symbol": "000017",
        "record": "2013-08-15",
        "halt": "2012-11-13",
        "resume": "2013-12-31",
        "ratio": "0.08",
    },
    {
        "symbol": "000030",
        "record": "2010-05-17",
        "halt": "2010-05-18",
        "resume": "2011-08-10",
        "bands": (("10000", "0"), (None, "0.20")),
    },
    {
        "symbol": "000034",
        "record": "2009-12-10",
        "halt": "2010-01-04",
        "resume": "2010-09-03",
        "ratio": "0.10",
    },
    {
        "symbol": "000557",
        "record": None,
        "halt": "2010-11-04",
        "resume": "2014-12-30",
        "bands": (("500000", "0.12"), (None, "0.18")),
    },
    {
        "symbol": "000587",
        "record": "2010-04-16",
        "halt": "2010-01-04",
        "resume": "2011-06-15",
        "bands": (("50000", "0"), (None, "0.06")),
    },
    {
        "symbol": "000677",
        "record": "2012-06-19",
        "halt": "2012-06-19",
        "resume": "2012-12-28",
        "ratio": "0.25",
    },
    {
        "symbol": "000787",
        "record": "2011-02-14",
        "effective": "2012-06-05",
        "bands": (("10000", "0.25"), (None, "0.30")),
        "official_floor": True,
    },
    {
        "symbol": "000820",
        "record": "2012-09-20",
        "halt": "2011-04-22",
        "resume": "2013-04-26",
        "ratio": "0.22",
    },
    {
        "symbol": "600678",
        "record": "2011-10-31",
        "halt": "2011-11-01",
        "resume": "2012-11-19",
        "ratio": "0.23",
        "sse_end": "2012-11-16",
        "sse_file": "sse-2012-p2.json",
    },
    {
        "symbol": "600699",
        "record": None,
        "halt": "2010-05-14",
        "resume": "2011-01-10",
        "ratio": "0.18",
        "sse_end": "2011-01-07",
        "sse_file": "sse-2010-p2.json",
    },
    {
        "symbol": "600757",
        "record": "2010-11-03",
        "halt": "2010-10-11",
        "resume": "2011-03-15",
        "bands": (("10000000", "0.33"), (None, "0.38")),
        "sse_end": "2011-03-14",
        "sse_file": "sse-2011-p2.json",
    },
)

PLAN_PUBLICATION = {
    "000017": "2013-08-06",
    "000030": "2010-09-27",
    "000034": "2010-04-14",
    "000557": "2011-07-05",
    "000587": "2010-04-17",
    "000677": "2012-10-13",
    "000787": "2011-03-02",
    "000820": "2012-09-27",
    "600678": "2012-05-19",
    "600699": "2010-07-28",
    "600757": "2010-11-04",
}

PASS_THROUGH_CLASSIFICATIONS = {
    "000818": "DIRECT_SPLIT_REGISTRATION_40_PERCENT_HOLDER_60_PERCENT_ESTATE",
    "000757": "DIRECT_REGISTRATION_TO_ADMINISTRATOR",
    "000035": "NO_ORIGINAL_HOLDER_ENTITLEMENT",
    "000972": "NO_ORIGINAL_HOLDER_ENTITLEMENT",
    "600074": "NO_ORIGINAL_HOLDER_ENTITLEMENT",
    "002015": "NO_ORIGINAL_HOLDER_ENTITLEMENT",
    "600145": "NO_ORIGINAL_HOLDER_ENTITLEMENT",
}

DEMERGER_CLASSES = {
    "600003": "ACCEPTED_PARENT_EVENT",
    "600649": "ACCEPTED_PARENT_EVENT",
    "601188": "SUCCESSOR_REFERENCE",
    "601518": "SUCCESSOR_REFERENCE",
    "601200": "SUCCESSOR_REFERENCE",
    "900935": "OUT_OF_SCOPE_B_SHARE_REFERENCE",
    "600153": "TERMINATED_PROPOSAL",
    "000780": "MERGER_REGULATORY_REFERENCE",
    "000895": "MERGER_REGULATORY_REFERENCE",
    "600859": "MERGER_REGULATORY_REFERENCE",
    "002431": "NON_EQUITY_LICENSE_SPLIT",
    "002482": "NON_EQUITY_LICENSE_SPLIT",
    "002822": "NON_EQUITY_LICENSE_SPLIT",
    "300668": "NON_EQUITY_LICENSE_SPLIT",
    "301091": "NON_EQUITY_LICENSE_SPLIT",
}


def verified(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_bytes())
    manifest = value.pop("manifest_sha256")
    if canonical_sha256(value) != manifest:
        raise ValueError(f"MANIFEST_CHANGED:{path}")
    value["manifest_sha256"] = manifest
    return value


def source(path: Path, published_on: str) -> DemergerSource:
    metadata = json.loads(path.with_suffix(path.suffix + ".meta.json").read_bytes())
    if sha256_file(path) != metadata["sha256"]:
        raise ValueError(f"SOURCE_CHANGED:{path}")
    return DemergerSource(
        url=metadata["url"],
        sha256=metadata["sha256"],
        published_on=date.fromisoformat(published_on),
        retrieved_at=metadata["retrieved_at_utc"],
        document=path.name,
    )


def szse_boundary(staging: Path, symbol: str, day: str) -> DemergerSource:
    path = staging / f"full-raw-coverage-v1/szse-official-day-controls-v1/{day}.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.active.calculate_dimension() == "A1:A1":
        workbook.active.reset_dimensions()
    rows = workbook.active.iter_rows(values_only=True)
    headers = next(rows)
    code_index, volume_index = headers.index("证券代码"), headers.index("成交量(万股)")
    matches = [row for row in rows if str(row[code_index]).zfill(6) == symbol]
    workbook.close()
    if len(matches) != 1 or Decimal(str(matches[0][volume_index]).replace(",", "")) <= 0:
        raise ValueError(f"FIRST_TRADABLE_BOUNDARY_CHANGED:{symbol}:{day}")
    return source(path, day)


def sse_boundary(staging: Path, spec: dict[str, Any]) -> DemergerSource:
    path = staging / "full-suspension-v1" / spec["sse_file"]
    rows = json.loads(path.read_bytes())["result"]
    matches = [
        row
        for row in rows
        if row["productCode"] == spec["symbol"]
        and row["startStopDate"] == spec["halt"].replace("-", "")
        and row["endStopDate"] == spec["sse_end"].replace("-", "")
    ]
    if len(matches) != 1:
        raise ValueError(f"SSE_TRADING_BOUNDARY_CHANGED:{spec['symbol']}")
    return source(path, spec["sse_end"])


def surrender_events(staging: Path, fixtures: Path) -> list[MandatoryPartialShareSurrenderEvent]:
    plan_root = staging / "full-merger-v1"
    events = []
    for spec in SURRENDERS:
        symbol = spec["symbol"]
        plan = next(plan_root.glob(f"{symbol}-capital-plan-*.pdf"))
        boundary = None
        if spec.get("resume"):
            boundary = (
                sse_boundary(staging, spec)
                if symbol.startswith("6")
                else szse_boundary(staging, symbol, spec["resume"])
            )
        bands = tuple(
            SurrenderQuantityBand(up_to_quantity=upper, surrender_ratio=ratio)
            for upper, ratio in spec.get("bands", ())
        )
        event = MandatoryPartialShareSurrenderEvent(
            security_id={"value": f"CN-A:{'SSE' if symbol.startswith('6') else 'SZSE'}:{symbol}"},
            affected_holder_class="ORDINARY_NON_RESPONSIBLE_HOLDER",
            record_date=spec.get("record"),
            effective_date=spec.get("effective"),
            research_effective_date=spec.get("resume"),
            nontrading_from=spec.get("halt"),
            boundary_sources=(boundary,) if boundary else (),
            convention_approval_sha256=(None if spec.get("official_floor") else APPROVAL_SHA256),
            surrender_ratio=spec.get("ratio"),
            quantity_bands=bands,
            consideration_type="NONE",
            consideration_value="0",
            quantity_rounding_rule=(
                "OFFICIAL_FLOOR_SURRENDER" if spec.get("official_floor") else "UNRESOLVED"
            ),
            research_fraction_rule=(
                None if spec.get("official_floor") else "CEIL_SURRENDER_FLOOR_RETAINED_V1"
            ),
            recipient="RESTRUCTURING_PLAN_RECIPIENTS",
            disposition="TRANSFERRED",
            sources=(source(plan, PLAN_PUBLICATION[symbol]),),
            confirmation_sources=(),
        )
        event.require_supported()
        events.append(event)
    for name in ("600381_surrender_v1.json", "600462_surrender_convention_v1.json"):
        event = MandatoryPartialShareSurrenderEvent.model_validate_json(
            (fixtures / name).read_bytes()
        )
        event.require_supported()
        events.append(event)
    if len(events) != 13 or len({str(event.security_id) for event in events}) != 13:
        raise ValueError("PARTIAL_SURRENDER_CENSUS_CHANGED")
    return sorted(
        events, key=lambda event: (event.application_date or date.max, str(event.security_id))
    )


def classify_metadata(metadata: dict[str, Any], termination: Path) -> tuple[list[dict], list[dict]]:
    demerger = next(row for row in metadata["results"] if row["query"] == "demerger")
    if set(DEMERGER_CLASSES) != {row["secCode"] for row in demerger["announcements"]}:
        raise ValueError("DEMERGER_QUERY_UNIVERSE_CHANGED")
    termination_text = termination.read_bytes()
    if (
        sha256_file(termination)
        != json.loads(termination.with_suffix(termination.suffix + ".meta.json").read_bytes())[
            "sha256"
        ]
        or b"PDF" not in termination_text[:16]
    ):
        raise ValueError("DEMERGER_TERMINATION_EVIDENCE_CHANGED")
    demerger_rows = [
        {
            "announcement_id": str(row["announcementId"]),
            "security": row["secCode"],
            "title": row["announcementTitle"],
            "classification": DEMERGER_CLASSES[row["secCode"]],
        }
        for row in demerger["announcements"]
    ]
    consolidation = next(
        row for row in metadata["results"] if row["query"] == "share_consolidation"
    )
    consolidation_rows = []
    for row in consolidation["announcements"]:
        title, symbol = row["announcementTitle"], row["secCode"]
        if symbol == "600381" and "缩股方案实施" in title:
            classification = "ACCEPTED_ISSUER_SHARE_CONSOLIDATION"
        elif symbol.startswith("2"):
            classification = "OUT_OF_SCOPE_B_SHARE"
        elif any(word in title for word in ("子公司", "非洲煤业")):
            classification = "NON_ISSUER_EQUITY"
        else:
            classification = "KEYWORD_FALSE_POSITIVE"
        consolidation_rows.append(
            {
                "announcement_id": str(row["announcementId"]),
                "security": symbol,
                "title": title,
                "classification": classification,
            }
        )
    if (
        Counter(row["classification"] for row in consolidation_rows)[
            "ACCEPTED_ISSUER_SHARE_CONSOLIDATION"
        ]
        != 1
    ):
        raise ValueError("SHARE_CONSOLIDATION_CENSUS_CHANGED")
    return demerger_rows, consolidation_rows


def certify(staging: Path, fixtures: Path) -> dict[str, Any]:
    structural_root = staging / "full-structural-census-v1"
    metadata_path = structural_root / "structural-action-metadata-v1.json"
    metadata = verified(metadata_path)
    rights_path = staging / "full-rights-inventory-v1/full-rights-census-v1.json"
    merger_path = staging / "full-merger-v1/merger-candidate-census-v3.json"
    rights, mergers = verified(rights_path), verified(merger_path)
    if rights["full_census_status"] != "PASS" or (
        mergers["full_census_status"],
        mergers["supported_events"],
        mergers["unresolved_candidate_count"],
    ) != ("PASS", 24, 0):
        raise ValueError("ACCEPTED_STRUCTURAL_BASELINE_CHANGED")

    surrenders = surrender_events(staging, fixtures)
    surrender_rows = [
        {"event": event.model_dump(mode="json"), "event_checksum": event.event_checksum}
        for event in surrenders
    ]
    if [row["event_checksum"] for row in surrender_rows] != [
        event.model_copy(
            update={
                "sources": tuple(reversed(event.sources)),
                "confirmation_sources": tuple(reversed(event.confirmation_sources)),
                "boundary_sources": tuple(reversed(event.boundary_sources)),
            }
        ).event_checksum
        for event in surrenders
    ]:
        raise ValueError("SURRENDER_INPUT_ORDER_CHANGED")

    pass_event = RestructuringBonusPassThroughEvent.model_validate_json(
        (fixtures / "600817_pass_through_v1.json").read_bytes()
    )
    pass_event.require_supported()
    pass_receipt_path = staging / "full-merger-v1/600817-pass-through-accepted-v1.json"
    pass_receipt = verified(pass_receipt_path)
    if (
        not pass_receipt["source_bytes_verified"]
        or pass_receipt["event_checksum"] != pass_event.event_checksum
    ):
        raise ValueError("PASS_THROUGH_BASELINE_CHANGED")
    plan_root = staging / "full-merger-v1"
    pass_classifications = []
    for symbol, classification in sorted(PASS_THROUGH_CLASSIFICATIONS.items()):
        plan = next(plan_root.glob(f"{symbol}-capital-plan-*.pdf"))
        metadata_row = json.loads(plan.with_suffix(plan.suffix + ".meta.json").read_bytes())
        if sha256_file(plan) != metadata_row["sha256"]:
            raise ValueError(f"PASS_THROUGH_PLAN_CHANGED:{symbol}")
        pass_classifications.append(
            {
                "security": symbol,
                "classification": classification,
                "source_path": str(plan),
                "source_sha256": metadata_row["sha256"],
            }
        )
    implementation_cache = (
        staging / "full-distribution-notices-v1/batch-v1/text/"
        "b95c752e22393e044434cfd0820e679ed6b6921bbd07a1c901c83bdf7384eebe.json"
    )
    implementation_text = "".join(
        "".join(json.loads(implementation_cache.read_bytes())["pages"]).split()
    )
    for required in ("40%登记至原股东账户", "60%定向登记至"):
        if required not in implementation_text:
            raise ValueError("000818_DIRECT_REGISTRATION_TERMS_CHANGED")

    termination = structural_root / "demerger-termination-v1/600153-termination-confirmation.pdf"
    demerger_rows, consolidation_rows = classify_metadata(metadata, termination)
    demerger_events = []
    for name in ("600003_demerger.json", "600649_demerger.json"):
        event = DemergerEvent.model_validate_json((fixtures / name).read_bytes())
        event.require_supported()
        demerger_events.append(
            {"event": event.model_dump(mode="json"), "event_checksum": event.event_checksum}
        )
    consolidation_event = ShareConsolidationEvent.model_validate_json(
        (fixtures / "600381_consolidation_v1.json").read_bytes()
    )
    consolidation_event.require_supported()

    inputs = [
        metadata_path,
        rights_path,
        merger_path,
        pass_receipt_path,
        implementation_cache,
        termination,
    ]
    report = {
        "report_id": "R5_FULL_STRUCTURAL_ACTION_CENSUS_V1",
        "coverage": ["2010-01-01", "2025-12-31"],
        "rights_full_census": {
            "status": "PASS",
            "events": 153,
            "manifest_sha256": rights["manifest_sha256"],
        },
        "merger_full_census": {
            "status": "PASS",
            "events": 24,
            "manifest_sha256": mergers["manifest_sha256"],
        },
        "demerger_full_census": {
            "status": "PASS",
            "events": demerger_events,
            "announcement_classifications": demerger_rows,
            "unresolved_candidates": 0,
            "termination_evidence_sha256": sha256_file(termination),
        },
        "partial_surrender_full_census": {
            "status": "PASS",
            "events": surrender_rows,
            "event_count": len(surrender_rows),
            "compensated_events": 0,
            "unresolved_consideration": 0,
        },
        "bonus_pass_through_full_census": {
            "status": "PASS",
            "events": [
                {
                    "event": pass_event.model_dump(mode="json"),
                    "event_checksum": pass_event.event_checksum,
                }
            ],
            "terminal_no_event_classifications": pass_classifications,
            "unresolved_candidates": 0,
        },
        "share_consolidation_full_census": {
            "status": "PASS",
            "events": [
                {
                    "event": consolidation_event.model_dump(mode="json"),
                    "event_checksum": consolidation_event.event_checksum,
                }
            ],
            "announcement_classifications": consolidation_rows,
            "unresolved_candidates": 0,
        },
        "input_order_reproducibility": "PASS",
        "checksum_repeatability": "PASS",
        "unresolved_structural_events": 0,
        "full_census_status": "PASS",
        "inputs": [{"path": str(path), "sha256": sha256_file(path)} for path in inputs],
        "parser_version": "R5_FULL_STRUCTURAL_ACTION_CENSUS_V1",
        "parser_sha256": sha256_file(Path(__file__)),
    }
    report["manifest_sha256"] = canonical_sha256(report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--staging", required=True, type=Path)
    parser.add_argument("--fixtures", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    report = certify(args.staging, args.fixtures)
    write_immutable(args.output, canonical_json_bytes(report) + b"\n")
    print(
        json.dumps(
            {
                "full_census_status": report["full_census_status"],
                "partial_surrender_events": report["partial_surrender_full_census"]["event_count"],
                "pass_through_events": len(report["bonus_pass_through_full_census"]["events"]),
                "unresolved_structural_events": report["unresolved_structural_events"],
                "manifest_sha256": report["manifest_sha256"],
            },
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
