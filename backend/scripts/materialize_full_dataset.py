"""Materialize the accepted 2010-2025 China A-share research dataset offline."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import warnings
from bisect import bisect_right
from collections import defaultdict
from datetime import UTC, datetime
from decimal import Context, Decimal, InvalidOperation, localcontext
from io import BytesIO
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from app.data.catalog import (
    DatasetPartition,
    finalize_manifest,
    load_manifest,
    verify_dataset,
    write_manifest,
)
from app.data.catalog.checksums import canonical_sha256, sha256_file
from app.data.catalog.freeze import require_unfrozen, write_immutable
from scripts.full_dataset_sources import (
    verify_source_contracts,
    verify_source_lineage,
    verify_source_manifest,
)
from scripts.publish_akshare_pilot import write_table
from scripts.reconcile_sse_cninfo import canonical, save, sha
from scripts.szse_st_recovery import risk_state

DATASET_ID = "cn-a-research-20100101-20251231-v1"
START = "2010-01-01"
END = "2025-12-31"
# Bump when raw parsing, units, ordering or serialization semantics change.
RAW_PARSER_VERSION = "OFFICIAL_RAW_INPUT_BOUND_V2"
RAW_FIELDS = (
    "security_id",
    "exchange",
    "symbol",
    "trading_date",
    "raw_open",
    "raw_high",
    "raw_low",
    "raw_close",
    "pre_close",
    "volume",
    "amount",
    "source_name",
    "risk_status",
    "source_lineage",
)


def decimal_text(value: object, *, scale: Decimal = Decimal(1)) -> str | None:
    if value in (None, "", "--", "-"):
        return None
    try:
        parsed = Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"INVALID_DECIMAL:{value}") from exc
    if not parsed.is_finite() or not scale.is_finite():
        raise ValueError(f"INVALID_DECIMAL:{value}")
    # Unit conversion must be exact and independent of the caller's Decimal context.
    with localcontext(
        Context(prec=max(28, len(parsed.as_tuple().digits) + len(scale.as_tuple().digits)))
    ):
        parsed *= scale
    return format(parsed, "f")


def validate_bar(row: dict[str, str | None]) -> bool:
    prices = [row[name] for name in ("raw_open", "raw_high", "raw_low", "raw_close")]
    if any(value is None for value in prices):
        return False
    opening, high, low, close = (Decimal(value) for value in prices if value is not None)
    volume = Decimal(row["volume"] or "-1")
    amount = Decimal(row["amount"] or "-1")
    if min(opening, high, low, close) <= 0 or volume < 0 or amount < 0:
        raise ValueError(f"INVALID_OHLCV:{row['security_id']}:{row['trading_date']}")
    if high < max(opening, close) or low > min(opening, close) or high < low:
        raise ValueError(f"INVALID_OHLC:{row['security_id']}:{row['trading_date']}")
    return True


def risk_from_name(name: object) -> str:
    return risk_state(str(name or ""))


class SecurityResolver:
    def __init__(self, master: list[dict[str, Any]]) -> None:
        self._intervals: dict[tuple[str, str], list[tuple[str, str, str]]] = defaultdict(list)
        for row in master:
            end = row["effective_to"] or END
            self._intervals[row["exchange"], row["symbol"]].append(
                (max(START, row["effective_from"]), min(END, end), row["security_id"])
            )
        for intervals in self._intervals.values():
            intervals.sort()

    def resolve(self, exchange: str, symbol: str, day: str) -> str | None:
        matches = [
            sid
            for start, end, sid in self._intervals.get((exchange, symbol), ())
            if start <= day <= end
        ]
        if len(matches) > 1:
            raise ValueError(f"AMBIGUOUS_SECURITY:{exchange}:{symbol}:{day}")
        return matches[0] if matches else None


class StatusResolver:
    def __init__(self, rows: list[dict[str, Any]]) -> None:
        self._rows: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
        for row in rows:
            self._rows[row["security_id"]].append(
                (row["effective_from"], row["effective_to"], row["status"])
            )
        for intervals in self._rows.values():
            intervals.sort()

    def resolve(self, security_id: str, day: str) -> str:
        matches = [status for start, end, status in self._rows[security_id] if start <= day <= end]
        if len(matches) != 1:
            raise ValueError(f"UNKNOWN_ST_STATUS:{security_id}:{day}")
        return matches[0]


class StreamingPartition:
    def __init__(self, path: Path, fields: tuple[str, ...], sort_keys: tuple[str, ...]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path
        self.fields = fields
        self.sort_keys = sort_keys
        self.schema = pa.schema(
            tuple(pa.field(field, pa.string(), nullable=True) for field in fields)
        )
        self.writer = pq.ParquetWriter(
            path,
            self.schema,
            compression="zstd",
            compression_level=3,
            use_dictionary=False,
            write_statistics=False,
            version="2.6",
            data_page_version="1.0",
            store_schema=True,
        )
        self.count = 0
        self.last_key: tuple[str, ...] | None = None
        self.digest = hashlib.sha256()
        self.digest.update(
            b'{"fields":'
            + json.dumps(fields, ensure_ascii=False, separators=(",", ":")).encode()
            + b',"rows":['
        )

    def write(self, rows: list[dict[str, str | None]]) -> None:
        if not rows:
            return
        for row in rows:
            key = tuple(row[name] or "" for name in self.sort_keys)
            if self.last_key is not None and key <= self.last_key:
                raise ValueError(f"NON_DETERMINISTIC_ROW_ORDER:{self.path}:{key}")
            self.last_key = key
            normalized = {field: row.get(field) for field in self.fields}
            if self.count:
                self.digest.update(b",")
            self.digest.update(
                json.dumps(
                    normalized,
                    ensure_ascii=False,
                    allow_nan=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ).encode()
            )
            self.count += 1
        self.writer.write_table(pa.Table.from_pylist(rows, schema=self.schema))

    def close(self, relative_path: str) -> DatasetPartition:
        self.writer.close()
        self.digest.update(b"]}")
        return DatasetPartition(
            relative_path=relative_path,
            media_type="application/vnd.apache.parquet",
            row_count=self.count,
            sha256=sha256_file(self.path),
            logical_sha256=self.digest.hexdigest(),
        )


def verified_bytes(
    path: Path, expected: tuple[str, str] | None = None
) -> tuple[bytes, dict[str, Any]]:
    raw = path.read_bytes()
    metadata_bytes = path.with_suffix(path.suffix + ".meta.json").read_bytes()
    metadata = json.loads(metadata_bytes)
    if sha(raw) != metadata["sha256"]:
        raise ValueError(f"SOURCE_CHECKSUM_CHANGED:{path}")
    if expected is not None and (sha(raw), sha(metadata_bytes)) != expected:
        raise ValueError(f"RAW_CHECKPOINT_INPUT_MISMATCH:{path}")
    return raw, metadata


def _raw_input_snapshot(
    exchange: str, paths: list[Path], resolver: SecurityResolver, status_resolver: StatusResolver
) -> tuple[str, dict[str, tuple[str, str]]]:
    sources = {}
    for path in paths:
        raw_sha = sha256_file(path)
        metadata_bytes = path.with_suffix(path.suffix + ".meta.json").read_bytes()
        if json.loads(metadata_bytes)["sha256"] != raw_sha:
            raise ValueError(f"SOURCE_CHECKSUM_CHANGED:{path}")
        sources[str(path)] = (raw_sha, sha(metadata_bytes))
    return canonical_sha256(
        {
            "parser_version": RAW_PARSER_VERSION,
            "exchange": exchange,
            "start": START,
            "end": END,
            "fields": RAW_FIELDS,
            "sources": sources,
            "security_intervals": sorted(resolver._intervals.items()),
            "status_intervals": sorted(status_resolver._rows.items()),
        }
    ), sources


def _require_raw_output(path: Path) -> None:
    if any(parent.is_symlink() for parent in (path, *path.parents)):
        raise ValueError(f"UNSAFE_RAW_CHECKPOINT_PATH:{path}")
    require_unfrozen(path)
    if path.is_dir():
        require_unfrozen(path / ".raw-write-check")


def raw_stage(
    build: Path,
    exchange: str,
    paths: list[Path],
    resolver: SecurityResolver,
    status_resolver: StatusResolver,
) -> tuple[list[DatasetPartition], list[dict[str, Any]], list[dict[str, str]]]:
    if exchange not in {"SSE", "SZSE"}:
        raise ValueError("INVALID_RAW_EXCHANGE")
    build = build.absolute()
    paths = sorted(path.absolute() for path in paths)
    if not paths or len({path.resolve() for path in paths}) != len(paths):
        raise ValueError("INVALID_RAW_SOURCE_SET")
    target = build / "raw" / exchange
    receipt_path = build / f"raw-{exchange.lower()}-checkpoint.json"
    temporary = build / f".raw-{exchange.lower()}.incomplete"
    for output in (target, receipt_path, receipt_path.with_suffix(".json.tmp"), temporary):
        _require_raw_output(output)
    if any(
        source.resolve().is_relative_to(build.resolve())
        for path in paths
        for source in (path, path.with_suffix(path.suffix + ".meta.json"))
    ):
        raise ValueError("RAW_SOURCE_OUTPUT_OVERLAP")
    if receipt_path.exists() != target.exists():
        raise ValueError(f"INCOMPLETE_RAW_CHECKPOINT:{exchange}")
    if receipt_path.exists() and (not receipt_path.is_file() or not target.is_dir()):
        raise ValueError(f"INVALID_RAW_CHECKPOINT:{exchange}")
    input_sha256, source_hashes = _raw_input_snapshot(exchange, paths, resolver, status_resolver)
    if receipt_path.exists():
        checkpoint = json.loads(receipt_path.read_bytes())
        if checkpoint.get("input_sha256") != input_sha256 or checkpoint.get("exchange") != exchange:
            raise ValueError(f"RAW_CHECKPOINT_INPUT_MISMATCH:{exchange}")
        recorded_sha = checkpoint.pop("checkpoint_sha256", None)
        if canonical_sha256(checkpoint) != recorded_sha:
            raise ValueError(f"INVALID_RAW_CHECKPOINT:{exchange}")
        parts = [DatasetPartition.model_validate(row) for row in checkpoint["partitions"]]
        relative_paths = [part.relative_path for part in parts]
        if relative_paths != sorted(set(relative_paths)) or any(
            not re.fullmatch(rf"raw/{exchange}/\d{{4}}-(0[1-9]|1[0-2])\.parquet", name)
            for name in relative_paths
        ):
            raise ValueError(f"INVALID_RAW_CHECKPOINT:{exchange}")
        children = list(target.iterdir())
        for child in children:
            _require_raw_output(child)
        if {build / name for name in relative_paths} != set(children) or any(
            not child.is_file() for child in children
        ):
            raise ValueError(f"INVALID_RAW_CHECKPOINT:{exchange}")
        if not all(sha256_file(build / part.relative_path) == part.sha256 for part in parts):
            raise ValueError(f"INVALID_RAW_CHECKPOINT:{exchange}")
        if _raw_input_snapshot(exchange, paths, resolver, status_resolver)[0] != input_sha256:
            raise ValueError(f"RAW_CHECKPOINT_INPUT_MISMATCH:{exchange}")
        return parts, checkpoint["lineage"], checkpoint["rejected"]

    if temporary.exists():
        for child in temporary.rglob("*"):
            _require_raw_output(child)
        shutil.rmtree(temporary)
    temporary.mkdir(parents=True)
    writers: dict[str, StreamingPartition] = {}
    buffers: dict[str, list[dict[str, str | None]]] = defaultdict(list)
    lineage: list[dict[str, Any]] = []
    rejected: list[dict[str, str]] = []
    flush_at = 2_000 if exchange == "SSE" else 20_000

    def flush(month: str) -> None:
        rows = buffers[month]
        if not rows:
            return
        if month not in writers:
            writers[month] = StreamingPartition(
                temporary / f"{month}.parquet",
                RAW_FIELDS,
                ("symbol", "trading_date", "security_id")
                if exchange == "SSE"
                else ("trading_date", "security_id"),
            )
        writers[month].write(rows)
        rows.clear()

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        for path in paths:
            raw, meta = verified_bytes(path, source_hashes[str(path)])
            lineage.append(
                {
                    "source_path": str(path),
                    "exchange": exchange,
                    "source_sha256": meta["sha256"],
                    "source_url": meta["url"],
                    "retrieved_at_utc": meta["retrieved_at_utc"],
                }
            )
            if exchange == "SSE":
                payload = json.loads(raw, parse_float=Decimal)
                symbol = path.stem
                previous = ""
                for values in payload["kline"]:
                    text = str(values[0])
                    day = f"{text[:4]}-{text[4:6]}-{text[6:]}"
                    if not START <= day <= END:
                        continue
                    security_id = resolver.resolve(exchange, symbol, day)
                    if security_id is None:
                        continue
                    if day <= previous:
                        raise ValueError(f"NON_MONOTONIC_SOURCE:{path}:{day}")
                    previous = day
                    row = {
                        "security_id": security_id,
                        "exchange": exchange,
                        "symbol": symbol,
                        "trading_date": day,
                        "raw_open": decimal_text(values[1]),
                        "raw_high": decimal_text(values[2]),
                        "raw_low": decimal_text(values[3]),
                        "raw_close": decimal_text(values[4]),
                        "pre_close": None,
                        "volume": decimal_text(values[5]),
                        "amount": decimal_text(values[6]),
                        "source_name": None,
                        "risk_status": status_resolver.resolve(security_id, day),
                        "source_lineage": "SSE_OFFICIAL_DAILY",
                    }
                    if validate_bar(row):
                        month = day[:7]
                        buffers[month].append(row)
                        if len(buffers[month]) >= flush_at:
                            flush(month)
            else:
                workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
                sheet = workbook.active
                if sheet.calculate_dimension() == "A1:A1":
                    sheet.reset_dimensions()
                sheet_rows = sheet.iter_rows(values_only=True)
                headers = next(sheet_rows)
                expected = (
                    "交易日期",
                    "证券代码",
                    "证券简称",
                    "前收",
                    "开盘",
                    "最高",
                    "最低",
                    "今收",
                    "涨跌幅\uff08%\uff09",
                    "成交量(万股)",
                    "成交金额(万元)",
                    "市盈率",
                )
                if tuple(headers) != expected:
                    raise ValueError(f"SZSE_HEADER_CHANGED:{path}")
                day_rows: list[dict[str, str | None]] = []
                for values in sheet_rows:
                    if not values or values[0] is None or values[1] is None:
                        continue
                    day, symbol = str(values[0])[:10], str(values[1]).zfill(6)
                    if not START <= day <= END:
                        continue
                    security_id = resolver.resolve(exchange, symbol, day)
                    if security_id is None:
                        continue
                    row = {
                        "security_id": security_id,
                        "exchange": exchange,
                        "symbol": symbol,
                        "trading_date": day,
                        "raw_open": decimal_text(values[4]),
                        "raw_high": decimal_text(values[5]),
                        "raw_low": decimal_text(values[6]),
                        "raw_close": decimal_text(values[7]),
                        "pre_close": decimal_text(values[3]),
                        "volume": decimal_text(values[9], scale=Decimal(10_000)),
                        "amount": decimal_text(values[10], scale=Decimal(10_000)),
                        "source_name": str(values[2]),
                        "risk_status": status_resolver.resolve(security_id, day),
                        "source_lineage": "SZSE_OFFICIAL_DAILY",
                    }
                    if validate_bar(row):
                        day_rows.append(row)
                    else:
                        rejected.append(
                            {
                                "security_id": security_id,
                                "trading_date": day,
                                "reason": "INCOMPLETE_OFFICIAL_DAILY_ROW_NOT_A_BAR",
                            }
                        )
                workbook.close()
                day_rows.sort(key=lambda row: row["security_id"] or "")
                if day_rows:
                    month = day_rows[0]["trading_date"][:7]  # type: ignore[index]
                    buffers[month].extend(day_rows)
                    if len(buffers[month]) >= flush_at:
                        flush(month)
    for month in sorted(buffers):
        flush(month)
    parts = []
    for month, writer in sorted(writers.items()):
        relative = f"raw/{exchange}/{month}.parquet"
        parts.append(writer.close(relative))
    if _raw_input_snapshot(exchange, paths, resolver, status_resolver)[0] != input_sha256:
        raise ValueError(f"RAW_CHECKPOINT_INPUT_MISMATCH:{exchange}")
    for output in (target, receipt_path, receipt_path.with_suffix(".json.tmp"), temporary):
        _require_raw_output(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary.rename(target)
    checkpoint = {
        "exchange": exchange,
        "input_sha256": input_sha256,
        "partitions": [part.model_dump(mode="json") for part in parts],
        "lineage": lineage,
        "rejected": rejected,
    }
    checkpoint["checkpoint_sha256"] = canonical_sha256(checkpoint)
    write_immutable(receipt_path, canonical(checkpoint))
    return parts, lineage, rejected


def adjustment_rows(inventory: dict[str, Any], authority: dict[str, Any]) -> list[dict[str, str]]:
    overlays: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in authority["official_cash_overlay_events"]:
        overlays[row["security_id"]].append(row)
    result = []
    for interval in sorted(
        inventory["intervals"], key=lambda row: (row["security_id"], row["effective_from"])
    ):
        variants = list(interval["in_interval_variants"].values())
        if len(variants) != 1:
            raise ValueError(f"AMBIGUOUS_ADJUSTMENT_CURVE:{interval['security_id']}")
        variant = variants[0]
        provider = {variant["baseline"][0]: Decimal(variant["baseline"][1])}
        provider.update((day, Decimal(value)) for day, value in variant["period_rows"])
        applicable = {
            row["effective_date"]: Decimal(row["factor_step"])
            for row in overlays[interval["security_id"]]
            if interval["effective_from"] <= row["effective_date"] <= interval["effective_to"]
        }
        dates = sorted(set(provider) | set(applicable))
        cumulative = Decimal(1)
        current: Decimal | None = None
        for day in dates:
            if day in provider:
                current = provider[day]
            if day in applicable:
                cumulative *= applicable[day]
            if current is None:
                raise ValueError(f"MISSING_ADJUSTMENT_BASELINE:{interval['security_id']}:{day}")
            result.append(
                {
                    "security_id": interval["security_id"],
                    "interval_from": interval["effective_from"],
                    "interval_to": interval["effective_to"],
                    "effective_date": day,
                    "hfq_factor": format(current * cumulative, "f"),
                    "lineage": "OFFICIAL_CASH_OVERLAY"
                    if day in applicable
                    else "SINA_HFQ_OFFICIAL_EVENT_BOUND",
                }
            )
    return result


def factor_at(rows: list[dict[str, str]], security_id: str, day: str) -> Decimal:
    scoped = [
        row
        for row in rows
        if row["security_id"] == security_id and row["interval_from"] <= day <= row["interval_to"]
    ]
    dates = [row["effective_date"] for row in scoped]
    index = bisect_right(dates, day) - 1
    if index < 0:
        raise ValueError(f"UNKNOWN_FACTOR:{security_id}:{day}")
    return Decimal(scoped[index]["hfq_factor"])


def action_row(
    event: dict[str, Any],
    event_checksum: str,
    *,
    event_kind: str,
    application_date: str,
    security_id: str | dict[str, str],
    acceptance: dict[str, Any],
) -> dict[str, str | None]:
    if isinstance(security_id, dict):
        security_id = security_id["value"]
    sources = event.get("sources") or event.get("allocation_sources") or ()
    first_source = sources[0] if sources else event.get("conversion_source") or {}
    return {
        "security_id": security_id,
        "event_kind": event_kind,
        "action_type": event.get("action_type", event_kind),
        "application_date": application_date,
        "effective_date": event.get("effective_date", application_date),
        "ex_date": event.get("ex_date"),
        "cash_per_share": event.get("cash_per_share"),
        "share_ratio": event.get("share_ratio"),
        "event_checksum": event_checksum,
        "policy_id": acceptance.get("policy_id", event.get("policy", event_kind)),
        "source_document_sha256": event.get("source_document_sha256") or first_source.get("sha256"),
        "source_url": event.get("source_url") or first_source.get("url"),
        "event_json": canonical(event).decode(),
        "acceptance_json": canonical(acceptance).decode(),
    }


def event_rows(path: Path, *, event_kind: str = "CORPORATE_ACTION") -> list[dict[str, str | None]]:
    result = []
    for item in json.loads(path.read_bytes())["events"]:
        event = item["event"]
        security = event["security_id"]
        acceptance = {key: value for key, value in item.items() if key != "event"}
        result.append(
            action_row(
                event,
                item["event_checksum"],
                event_kind=event_kind,
                application_date=event["effective_date"],
                security_id=security["value"] if isinstance(security, dict) else security,
                acceptance=acceptance,
            )
        )
    return result


def structural_event_rows(
    repository: Path, rights_path: Path, merger_path: Path, structural_path: Path
) -> list[dict[str, str | None]]:
    rows = []
    rights_report = json.loads(rights_path.read_bytes())
    rights_inventory = json.loads(Path(rights_report["inventory_path"]).read_bytes())
    for item in rights_inventory["rows"]:
        event = item["event"]
        rows.append(
            action_row(
                event,
                item["event_checksum"],
                event_kind="RIGHTS_ISSUE",
                application_date=event["effective_date"],
                security_id=event["security_id"]["value"],
                acceptance={key: value for key, value in item.items() if key != "event"},
            )
        )

    merger_census = json.loads(merger_path.read_bytes())
    for item in merger_census["executable_events"]:
        event = item["event"]
        rows.append(
            action_row(
                event,
                item["event_checksum"],
                event_kind="MERGER",
                application_date=event["effective_date"],
                security_id=event["parent_security"]["value"],
                acceptance={key: value for key, value in item.items() if key != "event"},
            )
        )
    if sum(row["event_kind"] == "MERGER" for row in rows) != 24:
        raise ValueError("MERGER_LEDGER_COUNT_CHANGED")

    structural = json.loads(structural_path.read_bytes())
    for section, event_kind, security_key, date_keys in (
        (
            "demerger_full_census",
            "DEMERGER",
            "parent_security",
            ("research_effective_date", "effective_date"),
        ),
        (
            "partial_surrender_full_census",
            "PARTIAL_SURRENDER",
            "security_id",
            ("research_effective_date", "effective_date"),
        ),
        (
            "bonus_pass_through_full_census",
            "BONUS_PASS_THROUGH",
            "security_id",
            ("effective_date",),
        ),
        (
            "share_consolidation_full_census",
            "SHARE_CONSOLIDATION",
            "security_id",
            ("effective_date",),
        ),
    ):
        for item in structural[section]["events"]:
            event = item["event"]
            security = event[security_key]
            application_date = next(event[key] for key in date_keys if event.get(key))
            rows.append(
                action_row(
                    event,
                    item["event_checksum"],
                    event_kind=event_kind,
                    application_date=application_date,
                    security_id=security["value"] if isinstance(security, dict) else security,
                    acceptance={"census": section, "status": "PASS"},
                )
            )
    return rows


def materialize(staging: Path, repository: Path) -> dict[str, Any]:
    state = json.loads((repository / ".research/project_state.json").read_bytes())
    preflight_path = repository / state["full_source_preflight_report"]
    preflight = json.loads(preflight_path.read_bytes())
    if (
        state["r5_status"] != "RUNNING_FULL_DATASET_MATERIALIZATION"
        or preflight["status"] != "PASS"
        or preflight["manifest_sha256"] != state["full_source_preflight_manifest_sha256"]
        or sha(preflight_path.read_bytes()) != state["full_source_preflight_file_sha256"]
    ):
        raise ValueError("FULL_DATASET_MATERIALIZATION_NOT_AUTHORIZED")
    verify_source_manifest(preflight_path, state["full_source_preflight_manifest_sha256"])
    master_path = staging / "full-master-public-v1-final/point-in-time-security-master.json"
    master_authority_path = repository / state["p0_public_report_path"]
    master = verify_source_manifest(
        master_path,
        state["p0_public_report_manifest_sha256"],
        master_authority_path=master_authority_path,
    )
    if (
        preflight["components"]["security_master"]["manifest_sha256"]
        != state["p0_public_report_manifest_sha256"]
    ):
        raise ValueError("SOURCE_PREFLIGHT_MASTER_REFERENCE_MISMATCH")
    final = staging.parent / DATASET_ID
    root = staging / "full-dataset-build-v1"
    build = root / "catalog-build"
    if final.exists():
        contracts = pq.read_table(final / "source-contracts.parquet").to_pylist()
        verify_source_contracts(contracts, repository, state)
        manifest = load_manifest(final / "manifest.json")
        verify_source_lineage(contracts, manifest.source_lineage)
        verify_dataset(final, manifest)
        retained_report: dict[str, Any] = json.loads(
            (root / "materialization-report.json").read_bytes()
        )
        return retained_report
    sse_authority_path = staging / "full-sse-st-v1/sse-st-authority-v4.json"
    sse_authority = json.loads(sse_authority_path.read_bytes())

    szse_st_path = staging / "full-szse-st-v1/dated-name-coverage-probe-v9.json"
    szse_st_authority = json.loads(szse_st_path.read_bytes())

    inventory_path = staging / "full-adjustment-audit-v1/retained-curve-inventory-v2.json"
    adjustment_authority_path = (
        staging / "full-adjustment-audit-v1/full-adjustment-authority-v2.json"
    )
    inventory = json.loads(inventory_path.read_bytes())
    adjustment_authority = json.loads(adjustment_authority_path.read_bytes())

    cash_path = staging / "full-distribution-notices-v1/cash-full-coverage-v43.json"
    bonus_path = staging / "full-bonus-inventory-v1/retained-bonus-ledger-v32.json"
    gap_distribution_path = (
        staging / "full-adjustment-audit-v1/accepted-gap-distribution-ledger-v1.json"
    )
    rights_path = staging / "full-rights-inventory-v1/full-rights-census-v1.json"
    merger_path = staging / "full-merger-v1/merger-candidate-census-v4.json"
    structural_path = staging / "full-structural-census-v1/structural-action-census-v1.json"

    calendar_path = staging / "full-calendar-v1/calendar-manifest-v3.json"
    calendar = json.loads(calendar_path.read_bytes())
    benchmark_root = staging / "full-benchmark-v1"
    benchmark_manifest = json.loads((benchmark_root / "acceptance-manifest-v2.json").read_bytes())

    source_contracts = [
        ("SOURCE_PREFLIGHT", preflight_path, preflight["manifest_sha256"], preflight["status"]),
        ("SECURITY_MASTER", master_path, state["p0_public_report_manifest_sha256"], "PASS"),
        (
            "SSE_ST",
            sse_authority_path,
            sse_authority["manifest_sha256"],
            sse_authority["SSE_HISTORICAL_ST_RECONSTRUCTION"],
        ),
        (
            "SZSE_ST",
            szse_st_path,
            state["full_szse_st_probe_sha256"],
            state["full_szse_st_status"],
        ),
        (
            "SUSPENSIONS",
            staging / "full-suspension-revalidation-v36/suspension-authority-v1.json",
            state["full_suspension_authority_manifest_sha256"],
            "PASS_PUBLIC_OFFICIAL_RECONSTRUCTION",
        ),
        (
            "ADJUSTMENTS",
            adjustment_authority_path,
            adjustment_authority["manifest_sha256"],
            adjustment_authority["adjustment_authority"],
        ),
        ("CASH_ACTIONS", cash_path, json.loads(cash_path.read_bytes())["manifest_sha256"], "PASS"),
        (
            "BONUS_ACTIONS",
            bonus_path,
            json.loads(bonus_path.read_bytes())["manifest_sha256"],
            "PASS",
        ),
        (
            "GAP_DISTRIBUTION_ACTIONS",
            gap_distribution_path,
            json.loads(gap_distribution_path.read_bytes())["manifest_sha256"],
            "PASS",
        ),
        (
            "RIGHTS_ACTIONS",
            rights_path,
            json.loads(rights_path.read_bytes())["manifest_sha256"],
            "PASS",
        ),
        (
            "MERGER_ACTIONS",
            merger_path,
            json.loads(merger_path.read_bytes())["manifest_sha256"],
            "PASS",
        ),
        (
            "STRUCTURAL_ACTIONS",
            structural_path,
            json.loads(structural_path.read_bytes())["manifest_sha256"],
            "PASS",
        ),
        ("CALENDAR", calendar_path, calendar["manifest_sha256"], "PASS"),
        (
            "BENCHMARK",
            benchmark_root / "acceptance-manifest-v2.json",
            benchmark_manifest["manifest_sha256"],
            benchmark_manifest["status"],
        ),
        (
            "MARKET_RULES_AND_COSTS",
            Path(state["full_transaction_cost_report"]),
            state["full_transaction_cost_report_manifest_sha256"],
            state["full_transaction_cost_authority"],
        ),
    ]
    contracts = [
        {
            "name": name,
            "source_path": str(path),
            "source_file_sha256": sha(path.read_bytes()),
            "source_manifest_sha256": manifest_sha,
            "status": status,
        }
        for name, path, manifest_sha, status in source_contracts
    ]
    verify_source_contracts(contracts, repository, state)

    build.mkdir(parents=True, exist_ok=True)

    resolver = SecurityResolver(master)
    sse_status = StatusResolver(sse_authority["intervals"])
    szse_status = StatusResolver(szse_st_authority["daily_replay"]["intervals"])
    raw_root = staging / "full-raw-coverage-v1"
    sse_parts, sse_lineage, sse_rejected = raw_stage(
        build,
        "SSE",
        [
            path
            for path in (raw_root / "sse-official-day-controls-v1").glob("*.json")
            if len(path.stem) == 6 and path.stem.isdigit()
        ],
        resolver,
        sse_status,
    )
    szse_parts, szse_lineage, szse_rejected = raw_stage(
        build,
        "SZSE",
        list((raw_root / "szse-official-day-controls-v1").glob("[0-9]*.xlsx")),
        resolver,
        szse_status,
    )

    factors = adjustment_rows(inventory, adjustment_authority)
    partial_path = raw_root / "szse-20110725-monthly-reconstruction-v6/report.json"
    partial = json.loads(partial_path.read_bytes())
    feature_rows = []
    for row in sorted(partial["results"], key=lambda item: item["security_id"]):
        raw_close = Decimal(row["official_target_close"])
        current = factor_at(factors, row["security_id"], partial["target_date"])
        feature_rows.append(
            {
                "security_id": row["security_id"],
                "trading_date": partial["target_date"],
                "raw_close": raw_close,
                "hfq_factor_at_day": current,
                "volume_lower": row["official_volume_interval_shares"][0],
                "volume_upper": row["official_volume_interval_shares"][1],
                "amount_lower": row["official_amount_interval_cny"][0],
                "amount_upper": row["official_amount_interval_cny"][1],
                "risk_status": szse_status.resolve(row["security_id"], partial["target_date"]),
                "lineage": "OFFICIAL_PARTIAL_FEATURE_ONLY_NON_TRADABLE",
            }
        )
    if len(feature_rows) != 239:
        raise ValueError("FEATURE_ONLY_ROW_COUNT_CHANGED")

    suspension_path = (
        staging / "full-suspension-revalidation-v36/normalized-suspension-events-v1.json"
    )
    suspension = json.loads(suspension_path.read_bytes())
    suspensions = []
    for exchange in ("sse", "szse"):
        for row in suspension[exchange]:
            suspensions.append(
                {
                    "security_id": row["security_id"],
                    "exchange": row["exchange"],
                    "suspended_at": row["suspended_at"],
                    "resumed_or_ended_at": row.get("resumed_or_ended_at", row.get("resumed_at")),
                    "execution_open_blocked": row["execution_open_blocked"],
                    "closure": row.get("closure"),
                    "source_files_json": canonical(row["source_files"]).decode(),
                    "source_sha256_json": canonical(row["source_sha256"]).decode(),
                }
            )

    corporate_actions = (
        event_rows(cash_path)
        + event_rows(bonus_path)
        + event_rows(gap_distribution_path)
        + structural_event_rows(repository, rights_path, merger_path, structural_path)
    )
    if len({row["event_checksum"] for row in corporate_actions}) != len(corporate_actions):
        raise ValueError("DUPLICATE_CORPORATE_ACTION")
    corporate_actions.sort(
        key=lambda row: (
            row["application_date"] or "",
            row["security_id"] or "",
            row["event_checksum"] or "",
        )
    )
    benchmark_source = benchmark_root / "benchmark.parquet"
    if sha256_file(benchmark_source) != benchmark_manifest["partition"]["sha256"]:
        raise ValueError("BENCHMARK_CHANGED")
    benchmark_target = build / "csi300.parquet"
    if not benchmark_target.exists():
        shutil.copyfile(benchmark_source, benchmark_target)

    lineage = sse_lineage + szse_lineage
    rejected = sse_rejected + szse_rejected
    costs = [
        {
            "effective_from": "2010-01-01",
            "effective_to": "2023-08-27",
            "commission_rate": "0.0003",
            "minimum_commission_cny": "5",
            "transfer_allowance_rate": "0.00001",
            "seller_stamp_duty_rate": "0.001",
            "scope": "APPROVED_RESEARCH_ACCOUNT_CONVENTION",
        },
        {
            "effective_from": "2023-08-28",
            "effective_to": END,
            "commission_rate": "0.0003",
            "minimum_commission_cny": "5",
            "transfer_allowance_rate": "0.00001",
            "seller_stamp_duty_rate": "0.0005",
            "scope": "APPROVED_RESEARCH_ACCOUNT_CONVENTION",
        },
    ]
    tables = (
        ("securities.parquet", master, ("security_id", "effective_from")),
        (
            "calendar.parquet",
            [{"session_date": day, "is_open": True} for day in calendar["sessions"]],
            ("session_date",),
        ),
        ("st-sse-intervals.parquet", sse_authority["intervals"], ("security_id", "effective_from")),
        (
            "st-szse-intervals.parquet",
            szse_st_authority["daily_replay"]["intervals"],
            ("security_id", "effective_from"),
        ),
        ("suspension-events.parquet", suspensions, ("security_id", "suspended_at")),
        ("adjustment-factors.parquet", factors, ("security_id", "interval_from", "effective_date")),
        ("feature-only-observations.parquet", feature_rows, ("security_id", "trading_date")),
        (
            "corporate-actions.parquet",
            corporate_actions,
            ("application_date", "security_id", "event_checksum"),
        ),
        ("source-lineage.parquet", lineage, ("exchange", "source_path")),
        ("rejected-observations.parquet", rejected, ("trading_date", "security_id")),
        ("source-contracts.parquet", contracts, ("name",)),
        ("transaction-costs.parquet", costs, ("effective_from",)),
    )
    partitions = [*sse_parts, *szse_parts]
    for relative, rows, keys in tables:
        partitions.append(write_table(build, relative, rows, keys))
    partitions.append(
        DatasetPartition.model_validate(benchmark_manifest["partition"]).model_copy(
            update={"relative_path": "csi300.parquet"}
        )
    )
    manifest = finalize_manifest(
        {
            "dataset_id": DATASET_ID,
            "provider": "OFFICIAL_SSE+OFFICIAL_SZSE+SINA+CSI",
            "provider_version": "R5_PUBLIC_OFFICIAL_RECONSTRUCTION_V1",
            "retrieval_timestamp": datetime.fromisoformat(
                preflight["approval"]["recorded_at_utc"].replace("Z", "+00:00")
            ).astimezone(UTC),
            "coverage_start": START,
            "coverage_end": END,
            "security_count": len({row["security_id"] for row in master}),
            "symbols": tuple(sorted({row["symbol"] for row in master})),
            "exchanges": ("SSE", "SZSE"),
            "fields": RAW_FIELDS,
            "adjustment_method": (
                "Separate event-bound HFQ factors plus three official cash overlays; "
                "raw execution only"
            ),
            "rule_version": state["market_rule_manifest_sha256"],
            "source_lineage": tuple(
                sorted(f"{row['name']}:{row['source_manifest_sha256']}" for row in contracts)
            ),
            "partitions": tuple(sorted(partitions, key=lambda part: part.relative_path)),
            "known_limitations": (
                "239 SZSE observations on 2011-07-25 are feature-only and never executable",
                "Provider adjustment factors normalize feature prices and never create "
                "portfolio entitlements",
                "Research-account transaction costs are an approved convention, not a "
                "universal broker claim",
            ),
        }
    )
    write_manifest(build / "manifest.json", manifest)
    verify_dataset(build, manifest)
    report = {
        "dataset_id": DATASET_ID,
        "coverage": [START, END],
        "securities": manifest.security_count,
        "security_intervals": len(master),
        "trading_sessions": len(calendar["sessions"]),
        "raw_partitions": len(sse_parts) + len(szse_parts),
        "raw_rows": sum(part.row_count for part in (*sse_parts, *szse_parts)),
        "sse_raw_rows": sum(part.row_count for part in sse_parts),
        "szse_raw_rows": sum(part.row_count for part in szse_parts),
        "feature_only_rows": len(feature_rows),
        "rejected_nonbar_rows": len(rejected),
        "adjustment_rows": len(factors),
        "corporate_action_rows": len(corporate_actions),
        "suspension_events": len(suspensions),
        "failed_sources": 0,
        "manifest_sha256": manifest.canonical_sha256,
        "manifest_file_sha256": sha((build / "manifest.json").read_bytes()),
        "dataset_audited": False,
        "dataset_frozen": False,
    }
    save(root / "materialization-report.json", report)
    build.rename(final)
    save(root / "publication-checkpoint.json", {**report, "published": True})
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--staging", required=True, type=Path)
    parser.add_argument("--repository", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(materialize(args.staging, args.repository)), flush=True)
