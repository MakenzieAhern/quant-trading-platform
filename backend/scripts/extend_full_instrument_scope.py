"""Extend accepted instrument evidence only for missing full-period candidates."""

import argparse
import json
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

from app.data.catalog.checksums import canonical_json_bytes, canonical_sha256, sha256_file
from app.data.catalog.freeze import write_immutable
from app.domain.markets.security import InstrumentClassification
from scripts.classify_pilot_instruments import classify, ipo_type
from scripts.reconcile_sse_cninfo import fetch


def validate_receipt(staging, member, receipt):
    if receipt["security_id"] != member["security_id"]:
        raise ValueError("RECEIPT_IDENTITY_CHANGED")
    if receipt["status"] == "UNRESOLVED":
        return receipt
    if receipt["status"] != "CLASSIFIED":
        raise ValueError("INVALID_RECEIPT_STATUS")
    result = receipt["result"]
    fact = InstrumentClassification.model_validate(result["classification"])
    if (
        fact.security_id.value != member["security_id"]
        or fact.effective_from.isoformat() != member["listing_date"]
    ):
        raise ValueError("RECEIPT_CLASSIFICATION_CHANGED")
    for source in result["sources"]:
        path = (staging / source["path"]).resolve()
        if not path.is_relative_to(staging.resolve()) or sha256_file(path) != source["sha256"]:
            raise ValueError("RECEIPT_SOURCE_CHANGED")
    return receipt


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("staging", "output", "freeze"):
        parser.add_argument("--" + name, type=Path, required=True)
    parser.add_argument("--recovery-attempt", type=int, choices=(0, 1, 2), default=0)
    parser.add_argument("--report", default="classification-manifest.json")
    args = parser.parse_args()
    staging, output = args.staging, args.output
    master_path = staging / "full-code-transitions-v1/master-code-intervals.json"
    master = json.loads(master_path.read_bytes())
    if (
        canonical_sha256({k: v for k, v in master.items() if k != "manifest_sha256"})
        != master["manifest_sha256"]
    ):
        raise ValueError("MASTER_MANIFEST_CHANGED")
    members = {
        r["security_id"]: r
        for r in master["records"]
        if r["listing_date"] <= "2025-12-31"
        and (not r["delisting_date"] or r["delisting_date"] >= "2010-01-01")
    }
    accepted_path = staging / "instrument-scope-20260912/approved-classifications.json"
    accepted_scope = json.loads(
        (accepted_path.parent / "approved-scope-manifest.json").read_bytes()
    )
    if sha256_file(accepted_path) != accepted_scope["classifications_sha256"]:
        raise ValueError("ACCEPTED_CLASSIFICATIONS_CHANGED")
    accepted = json.loads(accepted_path.read_bytes())
    supplement_path = staging / "pilot-v2-master/new-member-evidence.json"
    freeze = json.loads(args.freeze.read_bytes())
    if (
        sha256_file(supplement_path)
        != freeze["accepted_source_sha256"]["pilot-v2-master/new-member-evidence.json"]
    ):
        raise ValueError("FROZEN_SUPPLEMENT_CHANGED")
    accepted.extend(
        dict(
            classification=r["classification"], accepted_parent_sha256=sha256_file(supplement_path)
        )
        for r in json.loads(supplement_path.read_bytes())
    )
    reused = {}
    listing_conflicts = []
    for r in accepted:
        fact = InstrumentClassification.model_validate(r["classification"])
        sid = fact.security_id.value
        if sid in reused:
            raise ValueError("DUPLICATE_ACCEPTED_CLASSIFICATION")
        if sid in members:
            if fact.effective_from.isoformat() != members[sid]["listing_date"]:
                listing_conflicts.append(
                    dict(
                        security_id=sid,
                        accepted_listing_date=fact.effective_from.isoformat(),
                        candidate_listing_date=members[sid]["listing_date"],
                        status="REQUIRES_MASTER_RECONCILIATION",
                    )
                )
            reused[sid] = r

    snapshots = {}
    for folder in ("security-master-pilot", "full-master-evidence"):
        root = staging / folder
        manifest = json.loads((root / "source-snapshot-manifest.json").read_bytes())
        for meta in manifest["snapshots"]:
            path = root / "raw" / meta["raw_snapshot_filename"]
            snapshots[(folder, path.name)] = path, meta

    def source(folder, name):
        path, meta = snapshots[(folder, name)]
        if sha256_file(path) != meta["sha256"]:
            raise ValueError("REGISTER_SOURCE_CHANGED")
        return path, dict(
            path=str(path.relative_to(staging)),
            source_url=meta["source_url"],
            sha256=meta["sha256"],
            retrieval_timestamp_utc=meta["retrieved_at"],
        )

    sse = {}
    for folder, name in (
        ("security-master-pilot", "sse-current-1.json"),
        ("security-master-pilot", "sse-current-8.json"),
        ("full-master-evidence", "sse-a-boards-terminated.json"),
    ):
        path, meta = source(folder, name)
        for r in json.loads(path.read_bytes())["result"]:
            sse[r["A_STOCK_CODE"]] = r, meta
    path, szse_meta = source("security-master-pilot", "szse-current.xlsx")
    workbook = load_workbook(path, read_only=False, data_only=True)
    values = workbook.active.iter_rows(values_only=True)
    header = next(values)
    szse = {
        r["A股代码"]: r
        for row in values
        if (r := dict(zip(header, row, strict=True))).get("A股代码")
    }
    workbook.close()

    impact_path = staging / "pilot-v2-master/master-impact.json"
    if (
        sha256_file(impact_path)
        != freeze["accepted_source_sha256"]["pilot-v2-master/master-impact.json"]
    ):
        raise ValueError("FROZEN_MASTER_CORRECTION_CHANGED")
    corrections = {
        r["security_id"]: r for r in json.loads(impact_path.read_bytes())["changed_lifecycles"]
    }
    for conflict in listing_conflicts:
        sid = conflict["security_id"]
        correction = corrections.get(sid)
        if not correction or (
            correction["before"]["listing_date"] != conflict["accepted_listing_date"]
            or correction["after"]["listing_date"] != conflict["candidate_listing_date"]
        ):
            continue
        member = members[sid]
        raw, meta = sse[member["symbol"]]
        reused[sid] = classify(
            staging,
            dict(
                **member,
                instrument_type="COMMON_A_SHARE",
                sources=[meta],
                source_fields={k: raw[k] for k in ("A_STOCK_CODE", "LIST_DATE", "LIST_BOARD")},
            ),
        )
        conflict["status"] = "RESOLVED_BY_FROZEN_PILOT_V2_CORRECTION"
        conflict["accepted_correction_sha256"] = sha256_file(impact_path)

    def acquire(member):
        sid, code = member["security_id"], member["symbol"]
        checkpoint = output / "receipts" / (sid.replace(":", "-") + ".json")
        if checkpoint.exists():
            prior = validate_receipt(staging, member, json.loads(checkpoint.read_bytes()))
            if prior["status"] == "CLASSIFIED" or not args.recovery_attempt:
                return prior
            if not prior.get("error", "").startswith(("URLError:", "TimeoutError:")):
                return prior
            checkpoint = output / f"retry-{args.recovery_attempt}" / checkpoint.name
            if checkpoint.exists():
                return validate_receipt(staging, member, json.loads(checkpoint.read_bytes()))
        try:
            if member["board"] == "SSE_STAR":
                result = None
                for number in range(1, 7):
                    params = dict(
                        isPagination="true",
                        **{
                            "pageHelp.pageSize": 10,
                            "pageHelp.cacheSize": 1,
                            "pageHelp.pageNo": number,
                            "pageHelp.beginPage": number,
                            "pageHelp.endPage": number,
                        },
                        START_DATE=(
                            date.fromisoformat(member["listing_date"]) - timedelta(days=60)
                        ).isoformat(),
                        END_DATE=member["listing_date"],
                        SECURITY_CODE=code,
                        TITLE="",
                        BULLETIN_TYPE="",
                    )
                    url = (
                        "https://query.sse.com.cn/security/stock/queryCompanyBulletinNew.do?"
                        + urlencode(params)
                    )
                    name = f"ipo/{code}-p{number:03}.json"
                    data = json.loads(fetch(output, name, url, referer="https://www.sse.com.cn/"))
                    matches = [
                        r
                        for group in data["result"]
                        for r in group
                        if ipo_type(r, code, member["listing_date"]) is not None
                    ]
                    if matches:
                        kinds = {ipo_type(r, code, member["listing_date"]) for r in matches}
                        if len(kinds) != 1:
                            raise ValueError("CONFLICTING_IPO_TYPES")
                        row = min(matches, key=lambda r: (r["SSEDATE"], r["ORG_BULLETIN_ID"]))
                        meta = json.loads((output / (name + ".meta.json")).read_bytes())
                        fact = InstrumentClassification(
                            security_id={"value": sid},
                            instrument_type=next(iter(kinds)),
                            effective_from=member["listing_date"],
                            source_publication_date=row["SSEDATE"],
                            source_url=url,
                            source_document_sha256=meta["sha256"],
                            retrieval_timestamp=meta["retrieved_at_utc"],
                        )
                        result = dict(
                            classification=fact.model_dump(mode="json"),
                            ipo_notice_metadata=row,
                            sources=[
                                dict(
                                    path=str((output / name).relative_to(staging)),
                                    sha256=meta["sha256"],
                                    source_url=url,
                                    retrieval_timestamp_utc=meta["retrieved_at_utc"],
                                )
                            ],
                        )
                        break
                    if number >= (data["pageHelp"]["pageCount"] or 0):
                        break
                if result is None:
                    raise ValueError("DATED_IPO_TYPE_UNRESOLVED")
            else:
                if member["exchange"] == "SSE":
                    raw, meta = sse[code]
                    if raw["A_STOCK_CODE"] != code or raw["LIST_BOARD"] != "1":
                        raise ValueError("EXPLICIT_A_REGISTER_REQUIRED")
                    fields = {k: raw[k] for k in ("A_STOCK_CODE", "LIST_DATE", "LIST_BOARD")}
                elif code in szse:
                    raw, meta = szse[code], szse_meta
                    fields = {k: raw[k] for k in ("A股代码", "A股上市日期", "A股简称")}
                else:
                    old_key = ("security-master-pilot", f"szse-profile-{code}.json")
                    if old_key in snapshots:
                        path, meta = source(*old_key)
                        data = json.loads(path.read_bytes())
                    else:
                        url = (
                            "https://www.szse.cn/api/report/index/companyGeneralization?"
                            + urlencode(dict(secCode=code))
                        )
                        name = f"profiles/{code}.json"
                        if args.recovery_attempt:
                            name = f"recovery-{args.recovery_attempt}/" + name
                        data = json.loads(
                            fetch(
                                output,
                                name,
                                url,
                                referer=f"https://www.szse.cn/certificate/individual/index.html?code={code}",
                            )
                        )
                        receipt = json.loads((output / (name + ".meta.json")).read_bytes())
                        meta = dict(
                            path=str((output / name).relative_to(staging)),
                            source_url=url,
                            sha256=receipt["sha256"],
                            retrieval_timestamp_utc=receipt["retrieved_at_utc"],
                        )
                    if (
                        data.get("cols", {}).get("agdm") != "A股代码"
                        or data.get("data", {}).get("agdm") != code
                    ):
                        raise ValueError("EXPLICIT_A_REGISTER_REQUIRED")
                    fields = {k: data["data"][k] for k in ("agdm", "agssrq", "agjc")}
                result = classify(
                    staging,
                    dict(
                        **member,
                        instrument_type="COMMON_A_SHARE",
                        sources=[meta],
                        source_fields=fields,
                    ),
                )
            receipt = dict(security_id=sid, status="CLASSIFIED", result=result)
            validate_receipt(staging, member, receipt)
        except Exception as exc:
            receipt = dict(
                security_id=sid, status="UNRESOLVED", error=f"{type(exc).__name__}: {exc}"
            )
        write_immutable(checkpoint, canonical_json_bytes(receipt) + b"\n")
        print(sid, receipt["status"], receipt.get("error", ""), flush=True)
        return receipt

    pending = [members[k] for k in sorted(members.keys() - reused.keys())]
    with ThreadPoolExecutor(max_workers=2) as pool:
        receipts = list(pool.map(acquire, pending))
    combined = [*reused.values(), *(r["result"] for r in receipts if r["status"] == "CLASSIFIED")]
    combined.sort(key=lambda r: r["classification"]["security_id"]["value"])
    report = dict(
        parser_version="FULL_INSTRUMENT_EXTENSION_V1",
        parser_sha256=sha256_file(Path(__file__)),
        source_master_sha256=sha256_file(master_path),
        accepted_classifications_sha256=sha256_file(accepted_path),
        accepted_supplement_sha256=sha256_file(supplement_path),
        reused=len(reused),
        requested_candidates=len(members),
        classifications=combined,
        counts=dict(
            sorted(Counter(r["classification"]["instrument_type"] for r in combined).items())
        ),
        failures=[r for r in receipts if r["status"] != "CLASSIFIED"],
        listing_date_conflicts=listing_conflicts,
        full_master_completeness="NOT_CERTIFIED_BY_INSTRUMENT_CLASSIFICATION",
        full_dataset_created=False,
    )
    report["manifest_sha256"] = canonical_sha256(report)
    write_immutable(output / args.report, canonical_json_bytes(report) + b"\n")
    print(json.dumps({k: v for k, v in report.items() if k != "classifications"}))


if __name__ == "__main__":
    main()
