"""Reconcile all-board official lifecycle inventories against frozen Pilot v1."""

import argparse
import json
from collections import Counter
from datetime import date
from pathlib import Path

import pyarrow.parquet as pq
from openpyxl import load_workbook

from scripts.fetch_a_share_security_master import _row, _write_csv, _yearbook_baseline
from scripts.reconcile_sse_cninfo import canonical, save, sha
from scripts.sse_st_closure import anchor_names

START, END = "2023-01-01", "2024-07-31"
BOARDS = {"SSE_MAIN", "SSE_STAR", "SZSE_MAIN", "SZSE_CHINEXT"}
LIFECYCLE = ("listing_date", "delisting_date", "effective_from", "effective_to", "board")


def overlaps(row, start=START, end=END):
    return row["listing_date"] <= end and (
        not row["delisting_date"] or row["delisting_date"] >= start
    )


def require_board_coverage(rows, exit_rows):
    if {r["board"] for r in rows} != BOARDS:
        raise ValueError("MISSING_EXCHANGE_BOARD_SEGMENT")
    if {r["board"] for r in exit_rows} != BOARDS:
        raise ValueError("MISSING_EXIT_BOARD_SEGMENT")


def pilot_symbol(code, transitions, end=END):
    matches = [t for t in transitions if t["new_symbol"] == code]
    if len(matches) > 1:
        raise ValueError("AMBIGUOUS_CODE_TRANSITION")
    if not matches:
        return code, None
    transition = matches[0]
    if transition["effective_date"] <= end:
        raise ValueError("IN_PERIOD_CODE_TRANSITION_REQUIRES_INTERVAL_SPLIT")
    return transition["old_symbol"], transition


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--staging", type=Path, required=True)
    parser.add_argument("--v1", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--transitions", type=Path, required=True)
    args = parser.parse_args()
    roots = [args.staging / "security-master-pilot", args.staging / "full-master-evidence"]
    evidence = []

    def source(root, name):
        manifest = json.loads((root / "source-snapshot-manifest.json").read_bytes())
        entry = next(s for s in manifest["snapshots"] if s["raw_snapshot_filename"] == name)
        path = root / "raw" / name
        if sha(path.read_bytes()) != entry["sha256"]:
            raise ValueError("SOURCE_HASH_CHANGED:" + name)
        evidence.append({**entry, "path": path.relative_to(args.staging).as_posix()})
        return path

    transitions = json.loads(args.transitions.read_bytes())["transitions"]
    for transition in transitions:
        path = source(roots[1], transition["snapshot"])
        if (
            sha(path.read_bytes()) != transition["sha256"]
            or evidence[-1]["source_url"] != transition["source_url"]
        ):
            raise ValueError("CODE_TRANSITION_PROVENANCE_MISMATCH")

    def sse(name, root, board=None):
        payload = json.loads(source(root, name).read_bytes())
        raw = payload["result"]
        if len(raw) != int(payload["pageHelp"]["total"]):
            raise ValueError("INCOMPLETE_EXCHANGE_PAGE:" + name)
        result = []
        for item in raw:
            code = item["A_STOCK_CODE"]
            if not code or code == "-":
                continue  # Original A_STOCK_CODE column; B-only rows are not A shares.
            listed = date.fromisoformat(item["LIST_DATE"])
            exited = (
                date.fromisoformat(item["DELIST_DATE"])
                if item.get("DELIST_DATE") not in (None, "", "-")
                else None
            )
            result.append(
                _row(
                    "SSE",
                    code,
                    item["COMPANY_ABBR"],
                    board or ("SSE_STAR" if item["LIST_BOARD"] == "2" else "SSE_MAIN"),
                    listed,
                    exited,
                    evidence[-1]["source_url"],
                )
            )
        return result

    rows = sse("sse-current-1.json", roots[0], "SSE_MAIN")
    rows += sse("sse-current-8.json", roots[0], "SSE_STAR")
    exited = sse("sse-a-boards-terminated.json", roots[1])
    sheet = load_workbook(source(roots[0], "szse-current.xlsx"), data_only=True).active
    values = list(sheet.values)
    for values_row in values[1:]:
        item = dict(zip(values[0], values_row, strict=True))
        code = str(item.get("A股代码") or "")
        if not code:
            continue
        code, transition = pilot_symbol(code, transitions)
        if not code.startswith(("000", "001", "002", "003", "300", "301")):
            raise ValueError("UNCLASSIFIED_CURRENT_REGISTER_SYMBOL:" + code)
        rows.append(
            _row(
                "SZSE",
                code,
                transition["old_name"] if transition else item["A股简称"],
                transition["board"]
                if transition
                else ("SZSE_CHINEXT" if code.startswith(("300", "301")) else "SZSE_MAIN"),
                date.fromisoformat(str(item["A股上市日期"])),
                None,
                evidence[-1]["source_url"],
            )
        )
    sheet = load_workbook(source(roots[1], "szse-terminated.xlsx"), data_only=True).active
    values = list(sheet.values)
    excluded_b = []
    for code, name, listed, end in values[1:]:
        if code.startswith("200"):
            excluded_b.append(code)
            continue
        if not code.startswith(("000", "001", "002", "003", "300", "301")):
            raise ValueError("UNCLASSIFIED_EXIT_REGISTER_SYMBOL:" + code)
        exited.append(
            _row(
                "SZSE",
                code,
                name,
                "SZSE_CHINEXT" if code.startswith(("300", "301")) else "SZSE_MAIN",
                date.fromisoformat(listed),
                date.fromisoformat(end),
                evidence[-1]["source_url"],
            )
        )
    require_board_coverage(rows, exited)
    combined = {}
    conflicts = []
    for row in [*rows, *exited]:
        sid = row["security_id"]
        previous = combined.get(sid)
        if previous and any(previous[k] != row[k] for k in ("listing_date", "board")):
            conflicts.append({"security_id": sid, "before": previous, "after": row})
        combined[sid] = row
    relevant = sorted((r for r in combined.values() if overlaps(r)), key=lambda r: r["security_id"])
    anchor_root = args.staging / "sse-st-pilot/closure-20260910/anchor"
    anchor_pdf = anchor_root / "yearbook-2023.pdf"
    anchor_meta = json.loads(anchor_pdf.with_suffix(".pdf.meta.json").read_bytes())
    if sha(anchor_pdf.read_bytes()) != anchor_meta["sha256"]:
        raise ValueError("INDEPENDENT_SSE_BASELINE_CHANGED")
    all_anchor_codes = set(
        anchor_names(json.loads((anchor_root / "yearbook-2023-text.json").read_bytes()))
    )
    baseline_sse = {code for code in all_anchor_codes if code.startswith("6")}

    class CachedYearbook:
        def get(self, name, url, **kwargs):
            return source(roots[0], name).read_bytes()

    baseline_szse, _ = _yearbook_baseline(CachedYearbook())
    baseline_differences = {}
    for exchange, baseline in (("SSE", baseline_sse), ("SZSE", baseline_szse)):
        at_start = {
            r["symbol"]
            for r in relevant
            if r["exchange"] == exchange and r["listing_date"] <= "2022-12-31"
        }
        baseline_differences[exchange] = dict(
            missing=sorted(baseline - at_start), extra=sorted(at_start - baseline)
        )
    old = {
        r["security_id"]: r for r in pq.read_table(args.v1 / "candidate-master.parquet").to_pylist()
    }
    current = {r["security_id"]: r for r in relevant}
    added = sorted(current.keys() - old.keys())
    removed = sorted(old.keys() - current.keys())
    changed = [
        {"security_id": sid, "before": old[sid], "after": current[sid]}
        for sid in sorted(old.keys() & current.keys())
        if any(old[sid][k] != current[sid][k] for k in LIFECYCLE)
    ]
    args.output.mkdir(parents=True, exist_ok=True)
    for exchange in ("SSE", "SZSE"):
        _write_csv(
            args.output / f"{exchange.lower()}-security-master.csv",
            [r for r in relevant if r["exchange"] == exchange],
        )
    save(args.output / "candidate-master.json", relevant)
    save(
        args.output / "all-lifecycles.json",
        sorted(combined.values(), key=lambda r: r["security_id"]),
    )
    before = sorted(
        r["security_id"]
        for r in combined.values()
        if r["delisting_date"] and r["delisting_date"] < START and r["security_id"] not in old
    )
    report = dict(
        status="INVENTORY_COMPLETE_CLASSIFICATION_AND_INDEPENDENT_BASELINE_PENDING",
        v1_master_securities=len(old),
        v2_master_securities=len(relevant),
        new_active_during_pilot=added,
        new_exited_before_pilot=before,
        new_exited_during_pilot=[
            sid for sid in added if START <= current[sid]["delisting_date"] <= END
        ],
        new_exited_after_pilot=[sid for sid in added if current[sid]["delisting_date"] > END],
        removed=removed,
        changed_lifecycles=changed,
        conflicts=conflicts,
        boards=dict(sorted(Counter(r["board"] for r in relevant).items())),
        exit_inventory_boards=dict(sorted(Counter(r["board"] for r in exited).items())),
        future_ipo_excluded=sorted(
            r["security_id"] for r in combined.values() if r["listing_date"] > END
        ),
        prehistory_only=sorted(
            r["security_id"]
            for r in combined.values()
            if overlaps(r, "2022-12-01", "2022-12-31") and not overlaps(r)
        ),
        excluded_szse_b_share_codes=excluded_b,
        candidate_master_sha256=sha(canonical(relevant)),
        evidence=evidence,
        code_transitions=transitions,
        independent_baseline_differences=baseline_differences,
        independent_baseline_sse_b_shares=sorted(all_anchor_codes - baseline_sse),
        limitation=(
            "Candidate census, not ordinary-instrument/ST acceptance. "
            "No dataset publication or v1 modification."
        ),
    )
    save(args.output / "master-impact.json", report)
    print(
        json.dumps(
            {
                k: v
                for k, v in report.items()
                if k not in {"evidence", "future_ipo_excluded", "new_exited_before_pilot"}
            },
            ensure_ascii=True,
        )
    )
    if (
        conflicts
        or removed
        or any(d["missing"] or d["extra"] for d in baseline_differences.values())
    ):
        raise SystemExit(2)


if __name__ == "__main__":
    main()
